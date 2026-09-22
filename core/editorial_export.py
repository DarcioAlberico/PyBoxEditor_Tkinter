"""Exportação editorial semântica para HTML, EPUB3, DOCX e PDF pesquisável."""

from __future__ import annotations

import base64
import html
import io
import json
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

import fitz

from core.editorial_model import EditorialBlock, EditorialDocument


EXPORT_FORMATS = frozenset({"html", "epub", "docx", "pdf", "json", "txt"})
EXPORT_MODES = frozenset({"faithful", "clean", "hybrid"})


@dataclass(frozen=True)
class ExportOptions:
    format: str = "html"
    mode: str = "clean"
    include_diagnostics: bool = True
    include_review_report: bool = False
    source_pdf: str | Path | None = None
    title: str = ""

    def __post_init__(self) -> None:
        formato = str(self.format).casefold().lstrip(".")
        modo = str(self.mode).casefold()
        if formato not in EXPORT_FORMATS:
            raise ValueError(f"formato editorial não suportado: {formato!r}")
        if modo not in EXPORT_MODES:
            raise ValueError(f"modo editorial não suportado: {modo!r}")
        object.__setattr__(self, "format", formato)
        object.__setattr__(self, "mode", modo)


@dataclass(frozen=True)
class ExportReport:
    format: str
    files: tuple[str, ...]
    warnings: tuple[str, ...] = ()
    metadata: Mapping[str, Any] = field(default_factory=dict)


def _text(value: Any) -> str:
    if isinstance(value, Mapping):
        if "text" in value:
            return str(value["text"])
        if "fen" in value:
            return str(value["fen"])
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return "" if value is None else str(value)


def _fen(value: Any) -> str:
    return str(value.get("fen", "")) if isinstance(value, Mapping) else ""


#: Lado do tabuleiro desenhado quando o diagrama chega sem imagem, em pixels.
#: Oito casas de 24 px: 1 KB de PNG, e legível no tablet e no papel.
LADO_DO_DIAGRAMA_PX = 192


def imagem_do_diagrama(block: EditorialBlock) -> tuple[str, str]:
    """
    O PNG da figura em base64, e de onde ele veio — ``(imagem, origem)``.

    **Nenhum diagrama sai sem imagem** (item 3 da revisão de 2026-09-18). O
    caminho legado traz o recorte ou o desenho prontos, no bloco ou no valor;
    a Fase 4 não traz imagem nenhuma — o ``DiagramResult`` guarda a posição e
    os *hashes* do recorte, não os pixels —, e até aqui o `<figure>` dela saía
    com legenda e sem figura. Um FEN, porém, é tudo que o desenho precisa:
    `render_diagrama.desenhar` é o mesmo código que escreve o livro, e custa um
    quilobyte por diagrama.

    ``origem`` é ``"recorte"`` (o que veio da página), ``"desenho"`` (feito do
    FEN) ou ``"nenhuma"`` — e a última vira aviso no relatório da exportação.
    """
    value = block.decision.value
    candidatos = [block.metadata.get("png_base64")]
    if isinstance(value, Mapping):
        candidatos.append(value.get("png_base64"))
    for encoded in candidatos:
        if not encoded:
            continue
        try:
            base64.b64decode(str(encoded), validate=True)
        except (ValueError, TypeError):
            continue
        return str(encoded), "recorte"
    fen = _fen(value)
    if not fen:
        return "", "nenhuma"
    try:
        from core import lado_a_jogar as lado_jogar
        from core import render_diagrama
        png, _largura, _altura = render_diagrama.desenhar(
            fen, lado_px=LADO_DO_DIAGRAMA_PX,
            # O indicador de quem joga só entra quando o lado foi **lido**:
            # desenhá-lo a partir da convenção seria pôr no papel a afirmação
            # que o resto deste módulo existe para não fazer.
            lado_a_jogar=(lado_jogar.do_fen(fen)
                          if origem_do_lado(block) != "assumed" else None))
    except Exception:  # noqa: BLE001 — fonte ausente ou FEN torto não derruba a exportação
        return "", "nenhuma"
    return base64.b64encode(png).decode("ascii"), "desenho"


def origem_do_lado(block: EditorialBlock) -> str:
    """``legend``, ``manual``, ``explicit`` ou ``assumed`` — a convenção."""
    value = block.decision.value
    if isinstance(value, Mapping):
        origem = value.get("side_to_move_source")
        if origem:
            return str(origem)
    return str(block.metadata.get("side_to_move_source", "assumed"))


def legenda_do_diagrama(block: EditorialBlock) -> str:
    """
    O que a legenda do diagrama diz além do FEN: de quem é a vez, e se sobrou
    revisão.

    Os dois são **carimbo, e não diagnóstico**: saem em todos os modos,
    inclusive no limpo. O modo limpo tira a proveniência e o histórico de
    hipóteses, que são para quem revisa; um diagrama cuja posição ninguém
    conferiu, ou cujo lado a jogar é convenção, continua sendo isso no arquivo
    entregue ao leitor — esconder a ressalva é o que transformaria uma leitura
    de 94,5% de casas certas em afirmação.
    """
    from core import lado_a_jogar as lado_jogar

    fen = _fen(block.decision.value)
    partes = [lado_jogar.marca(lado_jogar.do_fen(fen), origem_do_lado(block))]
    if revisao_pendente(block):
        partes.append("não revisado")
    return "; ".join(partes)


def revisao_pendente(block: EditorialBlock) -> bool:
    """O diagrama ainda espera olho humano?"""
    value = block.decision.value
    estado = ""
    if isinstance(value, Mapping):
        estado = str(value.get("review_status", ""))
        if value.get("review_required"):
            return True
    estado = estado or str(block.metadata.get("review_status", ""))
    return (estado in {"review_required", "unresolved"}
            or block.decision.status == "unresolved")


def _audit(block: EditorialBlock, options: ExportOptions) -> str:
    if options.mode == "clean":
        return ""
    provenance = ", ".join(f"p{ref.page_index + 1}" for ref in block.source_refs)
    original = _text(block.decision.original_value)
    return (f'<details class="audit"><summary>Proveniência e auditoria</summary>'
            f'<span data-status="{html.escape(block.decision.status, quote=True)}">'
            f"Origem: {html.escape(provenance or 'desconhecida')}.</span>"
            f"<span>Hipótese original: {html.escape(original or _text(block.decision.value))}</span>"
            f"</details>")


def trechos_em_negrito(texto: str, block: EditorialBlock) -> list[tuple[str, bool]]:
    """
    O texto partido nos trechos em negrito e nos de fora deles.

    O adapter guarda o negrito como faixas de índice em `style["bold_spans"]`
    (é como `core/negrito.py` o mede, caractere a caractere), e os escritores
    daqui o ignoravam: o parágrafo saía inteiro em redondo, e a ênfase que o
    livro imprimiu sumia do HTML, do EPUB e do DOCX. Faixa fora do texto é
    cortada, e faixa vazia não vira `run` nenhum.
    """
    faixas = []
    for item in block.style.get("bold_spans") or []:
        try:
            inicio, fim = int(item[0]), int(item[1])
        except (TypeError, ValueError, IndexError):
            continue
        faixas.append((inicio, fim))
    saida: list[tuple[str, bool]] = []
    anterior = 0
    for inicio, fim in sorted(faixas):
        inicio, fim = max(inicio, anterior), min(fim, len(texto))
        if fim <= inicio:
            continue
        if inicio > anterior:
            saida.append((texto[anterior:inicio], False))
        saida.append((texto[inicio:fim], True))
        anterior = fim
    if anterior < len(texto):
        saida.append((texto[anterior:], False))
    return saida


def _com_negrito(texto: str, block: EditorialBlock) -> str:
    """O texto escapado, com `<strong>` onde o livro imprimiu negrito."""
    partes = trechos_em_negrito(texto, block)
    if not partes:
        return html.escape(texto)
    return "".join(f"<strong>{html.escape(trecho)}</strong>" if forte
                   else html.escape(trecho) for trecho, forte in partes)


def _block_html(block: EditorialBlock, options: ExportOptions) -> str:
    value = block.decision.value
    anchor = html.escape(block.id, quote=True)
    common = f' id="{anchor}" data-kind="{html.escape(block.kind, quote=True)}"'
    if block.kind == "heading":
        body = f"<h2{common}>{html.escape(_text(value))}</h2>"
    elif block.kind == "chess_sequence":
        body = (f'<div{common} class="chess-sequence" aria-label="Sequência de xadrez">'
                f"{html.escape(_text(value))}</div>")
    elif block.kind == "diagram":
        fen = _fen(value)
        data_fen = html.escape(fen, quote=True)
        ressalva = legenda_do_diagrama(block)
        alt = html.escape(str(block.metadata.get(
            "alt", f"Diagrama de xadrez: {fen} — {ressalva}")), quote=True)
        encoded, origem_da_imagem = imagem_do_diagrama(block)
        image = (f'<img alt="{alt}" src="data:image/png;base64,{encoded}">'
                 if encoded else "")
        body = (f'<figure{common} data-fen="{data_fen}" '
                f'data-side-source="{html.escape(origem_do_lado(block), quote=True)}" '
                f'data-image="{origem_da_imagem}"'
                + (' data-review="pending"' if revisao_pendente(block) else "")
                + ' aria-label="Diagrama de xadrez">'
                f"{image}<figcaption>{html.escape(fen or _text(value))}"
                f" <span class=\"ressalva\">({html.escape(ressalva)})</span>"
                f"</figcaption></figure>")
    elif block.kind in ("figure", "caption") and isinstance(value, Mapping):
        # A figura que **não** é tabuleiro: a faixa impressa acima do diagrama
        # e a página inteira que virou imagem (item 4 da revisão de
        # 2026-09-18). Saíam como `<figure data-fen="None">` com a legenda
        # cheia de JSON, porque toda `livro.Figura` chegava aqui como diagrama.
        encoded, origem_da_imagem = imagem_do_diagrama(block)
        rotulo = html.escape(str(value.get("warning") or "") or (
            "Cabeçalho do diagrama" if block.kind == "caption" else "Figura"),
            quote=True)
        corpo = (f'<img alt="{rotulo}" src="data:image/png;base64,{encoded}">'
                 if encoded else f"<figcaption>{rotulo}</figcaption>")
        body = f'<figure{common} data-image="{origem_da_imagem}">{corpo}</figure>'
    elif block.kind == "caption":
        body = f'<p{common} class="caption">{html.escape(_text(value))}</p>'
    elif block.kind == "table" and isinstance(value, Mapping):
        rows = value.get("rows", ())
        body = (f"<table{common}>" + "".join(
            "<tr>" + "".join(f"<td>{html.escape(_text(cell))}</td>" for cell in row) + "</tr>"
            for row in rows) + "</table>")
    elif block.kind == "page_break":
        body = f'<hr{common} class="page-break">'
    else:
        body = f"<p{common}>{_com_negrito(_text(value), block)}</p>"
    return body + _audit(block, options)


def _html_body(document: EditorialDocument, options: ExportOptions) -> str:
    pages = []
    for page in sorted(document.pages, key=lambda item: item.page_index):
        blocks = "\n".join(_block_html(block, options)
                             for block in sorted(page.blocks, key=lambda item: item.order))
        pages.append(f'<section id="page-{page.page_index + 1}" data-page="{page.page_index + 1}">{blocks}</section>')
    title = html.escape(options.title or document.title)
    language = html.escape(document.language or "und", quote=True)
    return (f'<!doctype html><html lang="{language}"><head><meta charset="utf-8">'
            f"<title>{title}</title><meta name=\"export-mode\" content=\"{options.mode}\">"
            "<link rel=\"stylesheet\" href=\"styles.css\"></head><body>"
            f'<main data-mode="{options.mode}"><h1>{title}</h1>{"".join(pages)}</main>'
            "</body></html>\n")


def _text_document(document: EditorialDocument) -> str:
    lines: list[str] = []
    for page in sorted(document.pages, key=lambda item: item.page_index):
        for block in sorted(page.blocks, key=lambda item: item.order):
            lines.append(_text(block.decision.value))
    return "\n\n".join(lines) + "\n"


def avisos_dos_diagramas(document: EditorialDocument) -> tuple[str, ...]:
    """
    Os diagramas que saíram sem imagem, e os que saíram sem revisão.

    O primeiro é defeito — uma figura vazia no livro entregue —, e o segundo é
    estado: o relatório da exportação é onde quem operou o livro descobre
    quantas posições ninguém conferiu, sem ter de abrir o arquivo.
    """
    sem_imagem: list[str] = []
    pendentes = 0
    for page in document.pages:
        for block in page.blocks:
            if block.kind != "diagram":
                continue
            if not imagem_do_diagrama(block)[0]:
                sem_imagem.append(block.id)
            pendentes += revisao_pendente(block)
    avisos = []
    if sem_imagem:
        avisos.append(f"{len(sem_imagem)} diagrama(s) sem imagem: "
                      + ", ".join(sem_imagem[:5])
                      + ("…" if len(sem_imagem) > 5 else ""))
    if pendentes:
        avisos.append(f"{pendentes} diagrama(s) exportados sem revisão")
    return tuple(avisos)


class EditorialExporter:
    """Exportador único: todos os destinos percorrem a mesma ordem de blocos."""

    def export(self, document: EditorialDocument, target: str | Path,
               options: ExportOptions | None = None) -> ExportReport:
        options = options or ExportOptions()
        caminho = Path(target)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        if options.format == "html":
            caminho.write_text(_html_body(document, options), encoding="utf-8")
            return ExportReport("html", (str(caminho),),
                                warnings=avisos_dos_diagramas(document),
                                metadata={"mode": options.mode})
        if options.format == "txt":
            caminho.write_text(_text_document(document), encoding="utf-8")
            return ExportReport("txt", (str(caminho),))
        if options.format == "json":
            document.save_json(caminho)
            return ExportReport("json", (str(caminho),), metadata={"schema": document.schema})
        if options.format == "epub":
            return self._epub(document, caminho, options)
        if options.format == "docx":
            return self._docx(document, caminho, options)
        return self._pdf(document, caminho, options)

    def _epub(self, document: EditorialDocument, target: Path,
               options: ExportOptions) -> ExportReport:
        body = _html_body(document, options).replace(
            '<html lang=', '<html xmlns="http://www.w3.org/1999/xhtml" lang=', 1)
        nav_links = "".join(
            f'<li><a href="content.xhtml#page-{page.page_index + 1}">Página {page.page_index + 1}</a></li>'
            for page in sorted(document.pages, key=lambda item: item.page_index))
        nav = (f'<?xml version="1.0" encoding="utf-8"?><!doctype html><html xmlns="http://www.w3.org/1999/xhtml" '
               'xmlns:epub="http://www.idpf.org/2007/ops" '
               f'lang="{html.escape(document.language)}"><head><title>{html.escape(document.title)}</title></head>'
               f"<body><nav epub:type=\"toc\" id=\"toc\"><ol>{nav_links}</ol></nav></body></html>")
        opf = (f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" '
               'version="3.0" unique-identifier="book-id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
               f'<dc:identifier id="book-id">{html.escape(document.document_id)}</dc:identifier>'
               f'<dc:title>{html.escape(document.title)}</dc:title><dc:language>{html.escape(document.language)}</dc:language>'
               '</metadata><manifest><item id="content" href="content.xhtml" media-type="application/xhtml+xml"/>'
               '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
               '<item id="css" href="styles.css" media-type="text/css"/></manifest>'
               '<spine><itemref idref="content"/></spine></package>')
        container = ('<?xml version="1.0" encoding="UTF-8"?><container version="1.0" '
                     'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                     '<rootfile full-path="OEBPS/package.opf" media-type="application/oebps-package+xml"/>'
                     '</rootfiles></container>')
        with zipfile.ZipFile(target, "w") as archive:
            archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
            archive.writestr("META-INF/container.xml", container)
            archive.writestr("OEBPS/content.xhtml", body)
            archive.writestr("OEBPS/nav.xhtml", nav)
            archive.writestr("OEBPS/package.opf", opf)
            archive.writestr(
                "OEBPS/styles.css",
                "body{font-family:serif} figure{break-inside:avoid}.chess-sequence{"
                "font-family:monospace}",
            )
        return ExportReport("epub", (str(target),),
                            warnings=avisos_dos_diagramas(document),
                            metadata={"mode": options.mode})

    def _docx(self, document: EditorialDocument, target: Path,
              options: ExportOptions) -> ExportReport:
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("DOCX requer a dependência opcional python-docx") from error
        word = Document()
        word.core_properties.title = options.title or document.title
        for page in sorted(document.pages, key=lambda item: item.page_index):
            for block in sorted(page.blocks, key=lambda item: item.order):
                text = _text(block.decision.value)
                if block.kind == "heading":
                    word.add_heading(text, level=1)
                elif block.kind == "chess_sequence":
                    paragraph = word.add_paragraph(style="Quote")
                    paragraph.add_run(text)
                elif block.kind == "diagram":
                    # A figura entra como **imagem**, e não como a linha de
                    # texto de antes: um DOCX de livro de xadrez sem tabuleiro
                    # nenhum é a mesma perda que o EPUB tinha (item 3 da
                    # revisão de 2026-09-18). O FEN e a ressalva ficam na
                    # legenda, que é onde se procura por eles.
                    encoded, _origem = imagem_do_diagrama(block)
                    if encoded:
                        from docx.shared import Pt as _Pt
                        word.add_picture(io.BytesIO(base64.b64decode(encoded)),
                                         width=_Pt(8 * 16))
                        word.paragraphs[-1].alignment = 1
                    paragraph = word.add_paragraph()
                    paragraph.add_run(
                        f"Diagrama de xadrez — FEN: {_fen(block.decision.value) or text}"
                        f" ({legenda_do_diagrama(block)})")
                elif (block.kind in ("figure", "caption")
                      and isinstance(block.decision.value, Mapping)):
                    encoded, _origem = imagem_do_diagrama(block)
                    if encoded:
                        word.add_picture(io.BytesIO(base64.b64decode(encoded)))
                        word.paragraphs[-1].alignment = 1
                    else:
                        word.add_paragraph(str(block.decision.value.get("warning")
                                               or "Figura"))
                elif block.kind == "table" and isinstance(block.decision.value, Mapping):
                    rows = list(block.decision.value.get("rows", ()) or ())
                    columns = max((len(row) for row in rows), default=0)
                    if columns:
                        table = word.add_table(rows=len(rows), cols=columns)
                        table.style = "Table Grid"
                        for row_index, row in enumerate(rows):
                            for column_index, cell in enumerate(row):
                                table.cell(row_index, column_index).text = _text(cell)
                else:
                    # O negrito do livro vira `run` em negrito, e não some: o
                    # adapter o guarda em `style["bold_spans"]` desde a Fase 1
                    # e este escritor o ignorava (item 4 da revisão de
                    # 2026-09-18).
                    paragraph = word.add_paragraph()
                    for trecho, forte in (trechos_em_negrito(text, block)
                                          or [(text, False)]):
                        paragraph.add_run(trecho).bold = forte or None
                if options.mode != "clean":
                    word.add_paragraph(f"[auditoria: {block.decision.status}]")
        word.save(target)
        return ExportReport("docx", (str(target),),
                            warnings=avisos_dos_diagramas(document),
                            metadata={"mode": options.mode})

    def _pdf(self, document: EditorialDocument, target: Path,
             options: ExportOptions) -> ExportReport:
        source = options.source_pdf or document.metadata.get("source_path")
        source_path = Path(source) if source else None
        if source_path and source_path.exists() and source_path.suffix.casefold() == ".pdf":
            pdf = fitz.open(source_path)
            while len(pdf) < len(document.pages):
                pdf.new_page()
        else:
            pdf = fitz.open()
            for _ in document.pages:
                pdf.new_page(width=612, height=792)
        text_items = 0
        failed_items = 0
        warnings: list[str] = []
        for page_model in sorted(document.pages, key=lambda item: item.page_index):
            if page_model.page_index >= len(pdf):
                # Página fora do PDF: contada como falha e anunciada — o
                # `continue` mudo deixava página em branco com "0 falhas".
                failed_items += len(page_model.blocks)
                warnings.append(f"página {page_model.page_index} fora do PDF "
                                f"({len(pdf)} páginas)")
                continue
            page = pdf[page_model.page_index]
            # As caixas do IR estão em **pixels do raster**; a página do PDF
            # está em pontos. A escala é a mesma de `ocr_export._escala_pdf`:
            # largura da página sobre largura da imagem, e `72/dpi` quando a
            # largura não foi gravada. Sem ela a caixa a 300 dpi caía 4,17×
            # fora da página.
            largura_imagem = page_model.metadata.get("image_width")
            dpi = float(page_model.metadata.get("dpi") or 300)
            escala = (page.rect.width / float(largura_imagem)
                      if largura_imagem else 72.0 / dpi)
            y = 40.0
            for block in sorted(page_model.blocks, key=lambda item: item.order):
                text = _text(block.decision.value)
                if block.kind == "diagram":
                    text = f"FEN: {_fen(block.decision.value) or text}"
                if not text:
                    continue
                ref = block.source_refs[0] if block.source_refs else None
                try:
                    if ref and ref.bbox and source_path:
                        x, y0, x2, y2 = (float(v) * escala for v in ref.bbox)
                        rect = fitz.Rect(x, y0, x2, y2) & page.rect
                        if rect.is_empty:
                            raise ValueError("caixa fora da página")
                        # O bloco é um parágrafo inteiro numa caixa: a camada
                        # tem de caber **dentro** dela, em várias linhas, senão
                        # a busca só acha o que sobrou na primeira. O corpo é o
                        # maior em que o texto cabe; `insert_textbox` devolve
                        # negativo quando não coube.
                        for tamanho in (14, 12, 11, 10, 9, 8, 7, 6, 5, 4, 3):
                            if page.insert_textbox(rect, text, fontsize=tamanho,
                                                   render_mode=3) >= 0:
                                break
                        else:
                            raise ValueError("texto não coube na caixa do bloco")
                        y = max(y, rect.y1 + 8)
                    else:
                        # Sem caixa a camada continua **invisível** (`render_mode=3`)
                        # sobre o scan: texto visível por cima da imagem original
                        # é o que o formato "pesquisável" existe para não fazer.
                        # Só o PDF criado do zero, sem imagem, recebe texto visível.
                        page.insert_text((40, y), text, fontsize=11,
                                         render_mode=3 if source_path else 0)
                        y += 18
                except Exception as error:  # noqa: BLE001 — contada, e o resto segue
                    failed_items += 1
                    warnings.append(f"bloco {block.id}: {type(error).__name__}: {error}")
                    continue
                text_items += 1
        pdf.save(target)
        pdf.close()
        return ExportReport("pdf", (str(target),), warnings=tuple(warnings), metadata={
            "pages": len(document.pages), "text_items": text_items,
            "failed_items": failed_items, "mode": options.mode,
        })

    @staticmethod
    def validate_output(path: str | Path, format: str) -> tuple[str, ...]:
        """Validação estrutural barata para CI e pré-visualização."""
        target = Path(path)
        formato = format.casefold().lstrip(".")
        errors: list[str] = []
        if not target.exists() or target.stat().st_size == 0:
            errors.append("arquivo de saída inexistente ou vazio")
        elif formato == "epub":
            try:
                with zipfile.ZipFile(target) as archive:
                    if archive.namelist()[0] != "mimetype":
                        errors.append("EPUB deve iniciar com mimetype")
                    if "OEBPS/content.xhtml" not in archive.namelist():
                        errors.append("EPUB sem content.xhtml")
            except zipfile.BadZipFile:
                errors.append("EPUB inválido")
        elif formato == "pdf":
            try:
                doc = fitz.open(target)
                if not len(doc):
                    errors.append("PDF sem páginas")
                doc.close()
            except Exception:
                errors.append("PDF inválido")
        return tuple(errors)
