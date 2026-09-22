"""Exportação dos resultados estruturados do OCR.

O módulo não substitui os exportadores históricos de livros. Ele oferece uma
fronteira simples para `PageResult`: texto, JSON, DOCX e PDF pesquisável com a
imagem/camada original preservada.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import fitz

from core.ocr_result import PageResult


@dataclass(frozen=True)
class ExportConfig:
    dpi: int = 300
    incluir_suspeitas: bool = True
    fonte_pdf: str = "helv"
    fonte_pdf_arquivo: str | None = None

    def __post_init__(self) -> None:
        if self.dpi <= 0:
            raise ValueError("dpi deve ser positivo")


def texto_da_pagina(pagina: PageResult) -> str:
    """Obtém texto preservando a estrutura disponível no resultado."""
    if pagina.text:
        return pagina.text
    linhas = [linha.text for linha in pagina.lines if linha.text]
    return "\n".join(linhas)


def exportar_texto(paginas: Sequence[PageResult], caminho: str | Path) -> Path:
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text("\n\n".join(texto_da_pagina(pagina) for pagina in paginas),
                       encoding="utf-8")
    return destino


def exportar_json(paginas: Sequence[PageResult], caminho: str | Path, *,
                  config: ExportConfig | None = None) -> Path:
    """Exporta resultados completos, opcionalmente omitindo suspeitas.

    Suspeitas pertencem ao diagnóstico e ficam no metadata da página para não
    contaminar o texto exportado. A opção explícita é útil para distribuir um
    JSON de produção sem expor a fila de revisão, sem alterar os objetos em
    memória nem apagar a trilha original.
    """
    config = config or ExportConfig()
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    paginas_json = []
    for pagina in paginas:
        dados_pagina = pagina.to_dict()
        if not config.incluir_suspeitas:
            metadata = dict(dados_pagina.get("metadata", {}))
            for chave in ("suspects", "suspeitas", "review_suspects"):
                metadata.pop(chave, None)
            dados_pagina["metadata"] = metadata
        paginas_json.append(dados_pagina)
    dados = {"pages": paginas_json}
    destino.write_text(json.dumps(dados, ensure_ascii=False, indent=2) + "\n",
                       encoding="utf-8")
    return destino


def _paragrafos(pagina: PageResult) -> list[str]:
    valor = pagina.metadata.get("paragraphs")
    if isinstance(valor, list):
        return [str(item) for item in valor if str(item).strip()]
    texto = texto_da_pagina(pagina)
    return [item.strip() for item in texto.split("\n\n") if item.strip()]


def exportar_docx(paginas: Sequence[PageResult], caminho: str | Path,
                  *, titulo: str = "OCR") -> Path:
    """Gera DOCX sem depender do PyMuPDF para o texto editável."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as erro:
        raise RuntimeError("DOCX exige python-docx; instale com .[docx]") from erro
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = titulo
    for indice, pagina in enumerate(paginas):
        if indice:
            doc.add_page_break()
        for regiao in pagina.regions:
            if regiao.type == "heading" and regiao.text:
                doc.add_heading(regiao.text, level=2)
        for paragrafo in _paragrafos(pagina):
            par = doc.add_paragraph(paragrafo)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.save(destino)
    return destino


def _escala_pdf(page: fitz.Page, pagina: PageResult, config: ExportConfig) -> float:
    largura = pagina.metadata.get("image_width")
    if largura:
        return page.rect.width / float(largura)
    return 72.0 / config.dpi


def _inserir_camada(page: fitz.Page, pagina: PageResult, config: ExportConfig) -> tuple[int, int]:
    escala = _escala_pdf(page, pagina, config)
    inseridos = 0
    falhos = 0
    # Palavra é a unidade correta para seleção e busca. Quando não há palavras,
    # a linha permite exportar resultados parciais sem fabricar boxes.
    itens = pagina.words or pagina.lines
    for item in itens:
        texto = getattr(item, "text", "")
        bbox = getattr(item, "bbox", None)
        if not texto or not bbox:
            continue
        x1, y1, x2, y2 = (valor * escala for valor in bbox)
        if x2 <= x1 or y2 <= y1:
            continue
        tamanho = max(1.0, min(72.0, (y2 - y1) * 0.78))
        try:
            argumentos = {"fontname": config.fonte_pdf, "fontsize": tamanho,
                          "render_mode": 3}
            if config.fonte_pdf_arquivo:
                argumentos["fontfile"] = config.fonte_pdf_arquivo
            page.insert_text((x1, y2), texto, **argumentos)
        except Exception:
            # Uma palavra com fonte/caractere incompatível não deve apagar as
            # demais da camada. O JSON continua contendo a informação completa.
            falhos += 1
            continue
        inseridos += 1
    return inseridos, falhos


def exportar_pdf_pesquisavel(input_pdf: str | Path, output_pdf: str | Path,
                             paginas: Sequence[PageResult], *,
                             config: ExportConfig | None = None) -> dict[str, int]:
    """Acrescenta camada invisível ao PDF original, sem rasterizá-lo."""
    config = config or ExportConfig()
    entrada = Path(input_pdf)
    saida = Path(output_pdf)
    if not entrada.exists():
        raise FileNotFoundError(str(entrada))
    saida.parent.mkdir(parents=True, exist_ok=True)
    documento = fitz.open(entrada)
    total_paginas = len(documento)
    inseridos = 0
    falhos = 0
    try:
        for indice, pagina in enumerate(paginas):
            if indice >= len(documento):
                break
            inseridos_pagina, falhos_pagina = _inserir_camada(documento[indice], pagina, config)
            inseridos += inseridos_pagina
            falhos += falhos_pagina
        documento.save(saida)
    finally:
        documento.close()
    return {"pages": min(len(paginas), total_paginas), "text_items": inseridos,
            "failed_items": falhos}
