"""
Testes da F59 — o diagrama como texto, com a fonte de xadrez dentro do arquivo.

O modo é opcional e o PNG continua padrão; o que estes testes prendem é que,
quando ele é pedido, o arquivo sai **completo**. Fonte embutida é o tipo de
recurso que falha pela metade: o EPUB abre, o tabuleiro aparece como
`rmblkans`, e nada no arquivo acusa o que faltou.

As quatro costuras do DOCX são o caso extremo disso — tipo de conteúdo, parte
ofuscada, entrada no `fontTable` e `w:embedTrueTypeFonts` —, e faltar qualquer
uma dá um arquivo que abre sem a fonte, ou não abre.

A geometria do EPUB (fila colada na fila, letra alinhada com a coluna) foi
medida no navegador durante a fase, e não aqui: exige motor de layout. O que
cabe na suíte é a estrutura, e é o que está abaixo.

**A F62 entra no mesmo arquivo porque é a mesma pergunta em outro lugar**: o
tabuleiro precisa da fonte de xadrez, e o texto corrido precisa de uma que
desenhe `♔♕♖♗♘♙`, `★`, `▼` — medido, a Times New Roman não tem uma figurina
sequer. Aquela é opção; esta entra sozinha, quando o texto pede.

Rodar sem pytest:      python tests/test_f59_fonte_embutida.py
"""

import os
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from core import exportar, livro
from core import render_diagrama as rd

FEN = "r1bqk2r/pp2bppp/2n1pn2/3p4/3P4/2N1PN2/PP2BPPP/R1BQK2R w - - 0 1"


def _figura(coordenadas=False, texto=True):
    """Uma figura de diagrama; `texto=False` imita o recorte, que não vira letra."""
    png, largura, altura = rd.desenhar(FEN, lado_px=128)
    fonte = rd.carregar()
    if not texto:
        return livro.Figura(png, largura, altura, origem="recorte")
    return livro.Figura(png, largura, altura, fen=FEN, origem="render",
                        linhas=rd.linhas(FEN, fonte), fonte=fonte.nome,
                        coordenadas=coordenadas)


def _paginas(*figuras):
    return [livro.PaginaExtraida(
        numero=0, blocos=[livro.Paragrafo("Prosa antes."), *figuras,
                          livro.Paragrafo("Prosa depois.")])]


def _epub(paginas, **kw):
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_epub(paginas, os.path.join(tmp, "livro.epub"), **kw)
    return zipfile.ZipFile(caminho)


def _docx(paginas, **kw):
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(paginas, os.path.join(tmp, "livro.docx"), **kw)
    return zipfile.ZipFile(caminho)


# ----------------------------------------------------------------------
# EPUB
# ----------------------------------------------------------------------

def test_o_epub_em_fonte_troca_a_imagem_por_texto():
    with _epub(_paginas(_figura()), diagramas="fonte") as z:
        nomes = z.namelist()
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

    assert "OEBPS/fonts/SkakNew-Diagram.otf" in nomes
    assert not [n for n in nomes if n.endswith(".png")], (
        "o diagrama saiu como texto, mas o PNG dele foi para o zip assim mesmo")
    assert '<div class="diagrama"' in pagina
    for linha in rd.linhas(FEN):
        assert f"<p>{linha}</p>" in pagina
    assert "<img" not in pagina


def test_o_epub_declara_a_fonte_no_manifesto_e_na_css():
    with _epub(_paginas(_figura()), diagramas="fonte") as z:
        opf = z.read("OEBPS/content.opf").decode("utf-8")
        css = z.read("OEBPS/estilo.css").decode("utf-8")

    assert 'media-type="font/otf"' in opf
    assert 'href="fonts/SkakNew-Diagram.otf"' in opf
    assert "@font-face" in css and 'font-family: "SkakNew-Diagram"' in css
    # Sem estes dois a fila descola da fila e o tabuleiro vira uma escada.
    assert "line-height: 1;" in css and "letter-spacing: 0;" in css


def test_o_apple_books_so_respeita_a_fonte_se_lhe_pedirem():
    """
    Sem o `ibooks:specified-fonts`, o Books troca a fonte embutida pela do
    leitor — e o tabuleiro vira `rmblkans` **só lá**, que é o pior tipo de
    defeito de formato: o arquivo passa em todos os outros leitores.
    """
    with _epub(_paginas(_figura()), diagramas="fonte") as z:
        opf = z.read("OEBPS/content.opf").decode("utf-8")

    assert 'property="ibooks:specified-fonts">true<' in opf
    assert "vocabulary.itunes.apple.com" in opf, "o prefixo do vocabulário faltou"


def test_o_recorte_continua_saindo_como_imagem():
    """
    Desenho e recorte convivem no mesmo livro. O recorte não tem letra que o
    descreva, então ele é imagem em qualquer modo.
    """
    with _epub(_paginas(_figura(), _figura(texto=False)), diagramas="fonte") as z:
        nomes = z.namelist()
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

    assert [n for n in nomes if n.endswith(".png")], "o recorte perdeu a imagem"
    assert '<div class="diagrama"' in pagina and "<img" in pagina


def test_o_modo_padrao_nao_embute_fonte_nenhuma():
    with _epub(_paginas(_figura())) as z:
        nomes = z.namelist()
        css = z.read("OEBPS/estilo.css").decode("utf-8")
        opf = z.read("OEBPS/content.opf").decode("utf-8")

    assert not [n for n in nomes if "fonts/" in n]
    assert "@font-face" not in css
    assert "ibooks" not in opf
    assert [n for n in nomes if n.endswith(".png")]


def test_as_coordenadas_saem_em_texto_no_epub():
    """
    A fonte não tem `a`–`h` nem `7` e `8` — as letras que sobrariam para rótulo
    desenham casa. Por isso o rótulo é um `<i>` de fonte de texto dentro de uma
    caixa que mede uma casa.
    """
    with _epub(_paginas(_figura(coordenadas=True)), diagramas="fonte") as z:
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")
        css = z.read("OEBPS/estilo.css").decode("utf-8")

    assert '<p class="colunas">' in pagina
    assert pagina.count('<span class="col">') == 8
    assert '<span class="rot"><i>8</i></span>' in pagina
    # a classe é o que impede o seletor das colunas de alargar o rótulo da fila
    assert "span.col {" in css and "span.rot {" in css


# ----------------------------------------------------------------------
# A ofuscação, que é o pedaço do DOCX que ninguém vê falhar
# ----------------------------------------------------------------------

def test_a_ofuscacao_e_reversivel_e_so_toca_os_32_primeiros_bytes():
    dados = bytes(range(256)) * 4
    ofuscado = exportar.ofuscar(dados)

    assert exportar.ofuscar(ofuscado) == dados, "o XOR não fechou"
    assert ofuscado[32:] == dados[32:], "mexeu além do cabeçalho"
    assert ofuscado[:32] != dados[:32]


def test_a_chave_sai_do_guid_em_ordem_inversa():
    """
    É a parte que se erra em silêncio: com a ordem trocada o arquivo continua
    sendo um zip válido, e só o Word reclama — na máquina de outra pessoa.
    """
    guid = "{01020304-0506-0708-090A-0B0C0D0E0F10}"
    esperado = bytes(range(1, 17))[::-1]
    saida = exportar.ofuscar(bytes(32), guid)
    assert saida[:16] == esperado
    assert saida[16:32] == esperado, "a chave não foi aplicada duas vezes"


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------

def test_o_docx_em_fonte_tem_as_quatro_costuras():
    with _docx(_paginas(_figura()), diagramas="fonte") as z:
        nomes = z.namelist()
        tipos = z.read("[Content_Types].xml").decode("utf-8")
        tabela = z.read("word/fontTable.xml").decode("utf-8")
        rels = z.read("word/_rels/fontTable.xml.rels").decode("utf-8")
        ajustes = z.read("word/settings.xml").decode("utf-8")
        embutida = z.read("word/fonts/fonte1.odttf")

    assert 'Extension="odttf"' in tipos and exportar.TIPO_OFUSCADO in tipos
    assert 'w:name="SkakNew-Diagram"' in tabela and "w:embedRegular" in tabela
    assert exportar.GUID_DA_FONTE in tabela
    assert 'Id="rIdFonte1"' in rels and "fonts/fonte1.odttf" in rels
    assert "<w:embedTrueTypeFonts/>" in ajustes
    assert "word/fonts/fonte1.odttf" in nomes

    with open(rd.carregar().arquivo, "rb") as f:
        assert exportar.ofuscar(embutida) == f.read(), (
            "a fonte embutida não é a fonte do disco")


def test_a_ordem_do_settings_e_a_do_esquema():
    """
    `CT_Settings` tem sequência fixa: `w:embedTrueTypeFonts` vem depois do
    `w:zoom`. Fora de lugar, o Word acusa arquivo corrompido — e o teste que
    olha só "está lá?" passaria.
    """
    with _docx(_paginas(_figura()), diagramas="fonte") as z:
        ajustes = z.read("word/settings.xml").decode("utf-8")

    assert ajustes.index("<w:embedTrueTypeFonts/>") > ajustes.index("<w:zoom")
    assert ajustes.index("<w:embedTrueTypeFonts/>") < ajustes.index("<w:defaultTabStop")


def test_o_tabuleiro_do_docx_e_texto_na_familia_certa():
    with _docx(_paginas(_figura()), diagramas="fonte") as z:
        documento = z.read("word/document.xml").decode("utf-8")

    for linha in rd.linhas(FEN):
        assert f"<w:t>{linha}</w:t>" in documento or linha in documento
    assert 'w:ascii="SkakNew-Diagram"' in documento
    assert 'w:hAnsi="SkakNew-Diagram"' in documento
    # entrelinha exata: no automático o Word abre uma faixa branca entre as filas
    assert 'w:lineRule="exact"' in documento


def test_no_docx_o_diagrama_com_coordenadas_sai_em_imagem():
    """
    Alinhar rótulo de outra fonte sobre as casas exigiria uma tabela de 81
    células por diagrama. No EPUB isso são três linhas de CSS; aqui, não — e a
    imagem já traz as coordenadas desenhadas.
    """
    with _docx(_paginas(_figura(coordenadas=True)), diagramas="fonte") as z:
        documento = z.read("word/document.xml").decode("utf-8")
        nomes = z.namelist()

    assert "<w:drawing>" in documento
    assert rd.linhas(FEN)[0] not in documento
    assert [n for n in nomes if n.startswith("word/media/")]
    assert not [n for n in nomes if n.startswith("word/fonts/")], (
        "não há texto na fonte, então não há por que embutir a fonte")


def test_o_docx_com_fonte_embutida_continua_sendo_um_docx():
    """
    **Abre com o `python-docx`, e não com `in`.** Os outros testes deste arquivo
    conferem por substring, e foi assim que um `<Default>` inserido depois da
    declaração XML — fora do `<Types>`, portanto num segundo elemento raiz —
    passou por eles: todas as substrings estavam lá, e o arquivo tinha deixado
    de ser XML. Quem parseia acusa; quem procura pedaço, não.
    """
    from docx import Document

    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(_paginas(_figura()),
                                 os.path.join(tmp, "livro.docx"),
                                 diagramas="fonte")
    doc = Document(caminho)
    linhas = [p.text for p in doc.paragraphs]
    for linha in rd.linhas(FEN):
        assert linha in linhas, "o tabuleiro não voltou como parágrafo"
    assert "Prosa antes." in linhas


def test_o_docx_padrao_nao_embute_nada():
    with _docx(_paginas(_figura())) as z:
        nomes = z.namelist()
        tipos = z.read("[Content_Types].xml").decode("utf-8")

    assert not [n for n in nomes if n.startswith("word/fonts/")]
    assert "odttf" not in tipos


# ----------------------------------------------------------------------
# A fonte dos símbolos do texto (F62)
# ----------------------------------------------------------------------

FIGURINAS = "♔♕♖♗♘♙"


def _pagina_de_texto(texto):
    return [livro.PaginaExtraida(numero=0, blocos=[livro.Paragrafo(texto)])]


def test_a_fonte_de_recurso_cobre_as_figurinas():
    """
    Medido no alfabeto do modelo: a Times New Roman não desenha **uma figurina
    sequer**, e a DejaVu Serif também não. Sem embutir, o livro exportado depende
    de o leitor ter uma fonte que as tenha.
    """
    escolha = exportar.fonte_dos_simbolos(f"1.e4 {FIGURINAS} 2.d4")
    assert escolha, "nenhuma fonte de recurso foi escolhida"
    familia, caminho, cobertos = escolha
    assert "Symbols" in familia
    assert os.path.exists(caminho)
    assert set(FIGURINAS) <= set(cobertos)


def test_sem_simbolo_no_texto_nao_entra_fonte_nenhuma():
    """641 KB é caro para um livro que não tem o que desenhar com eles."""
    assert exportar.fonte_dos_simbolos("Um texto so com letras e 1234.") is None
    with _epub(_pagina_de_texto("So letras aqui, sem figurina.")) as z:
        assert not [n for n in z.namelist() if "fonts/" in n]


def test_a_acentuacao_fica_com_a_fonte_do_texto():
    """
    O corte está em U+2000, e ele foi medido: abaixo dele estão `©`, `±`, `½` e
    as letras acentuadas, que qualquer fonte de texto desenha — e que ficariam
    de outra cor e outro peso numa fonte de símbolos.
    """
    escolha = exportar.fonte_dos_simbolos("café ± ½ © ² Č ♖")
    assert escolha and escolha[2] == "♖"


def test_o_simbolo_sai_embrulhado_e_a_letra_nao():
    with _epub(_pagina_de_texto("Depois de 20.♖xf3 ♕d5 as brancas estão ±.")) as z:
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")
        css = z.read("OEBPS/estilo.css").decode("utf-8")
        nomes = z.namelist()

    assert '<span class="sim">♖</span>xf3' in pagina
    assert '<span class="sim">♕</span>d5' in pagina
    assert "±" in pagina and '<span class="sim">±' not in pagina
    assert "span.sim {" in css and "@font-face" in css
    assert [n for n in nomes if n.startswith("OEBPS/fonts/")]


def test_simbolos_seguidos_vao_num_span_so():
    """Um `<span>` por caractere dobraria o XHTML sem mudar um pixel."""
    with _epub(_pagina_de_texto("As peças: ♔♕♖♗♘♙ e mais nada.")) as z:
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

    assert f'<span class="sim">{FIGURINAS}</span>' in pagina
    assert pagina.count('class="sim"') == 1


def test_no_docx_so_o_run_do_simbolo_troca_de_fonte():
    """
    No DOCX não há `unicode-range`: a fonte é atributo do run, então o parágrafo
    tem de ser partido onde a família muda.
    """
    from docx import Document

    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(_pagina_de_texto("Depois de 20.♖xf3 as brancas."),
                                 os.path.join(tmp, "livro.docx"))
    doc = Document(caminho)
    runs = doc.paragraphs[0].runs

    assert len(runs) == 3, [r.text for r in runs]
    assert runs[0].text == "Depois de 20." and runs[0].font.name is None
    assert runs[1].text == "♖" and "Symbols" in (runs[1].font.name or "")
    assert runs[2].text.startswith("xf3")


def test_o_docx_embute_a_fonte_dos_simbolos():
    with _docx(_pagina_de_texto("Depois de 20.♖xf3 as brancas.")) as z:
        nomes = z.namelist()
        tabela = z.read("word/fontTable.xml").decode("utf-8")

    assert [n for n in nomes if n.startswith("word/fonts/")]
    assert "Symbols" in tabela and "w:embedRegular" in tabela


def test_modo_de_diagrama_invalido_reclama_nos_dois_formatos():
    with tempfile.TemporaryDirectory() as tmp:
        for escrever, ext in ((exportar.para_epub, "epub"),
                              (exportar.para_docx, "docx")):
            with pytest.raises(ValueError) as erro:
                escrever(_paginas(_figura()), os.path.join(tmp, f"x.{ext}"),
                         diagramas="svg")
            assert "svg" in str(erro.value)


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
