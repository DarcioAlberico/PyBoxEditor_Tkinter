"""
PD-03 (`docs/ROADMAP_PENDENCIAS.md`) — a célula de tabela guarda a quebra de linha.

A F72 registrou: `W: Win (1 ♖e1!)` e `B: Draw (1...♖a2!)` são duas linhas no
papel e saíam numa, separadas por espaço. Para a tabela do Nunn isso não perde
informação; para uma de frases independentes, perderia. A célula passa a levar
`\\n` entre as linhas, cada saída o escreve do seu jeito, e o texto corrido da
página — de onde saem o alfabeto dos símbolos e a régua do corpus — continua
com espaço.
"""

import sys
import zipfile

import cv2
import numpy as np
import pytest
from PIL import Image

from core import exportar, livro
from core.livro import PaginaExtraida, Paragrafo, Tabela
from core.services.box_service import BoxService

CELULA = "W: Win (1 ♖e1!)\nB: Draw (1...♖a2!)"


def _xhtml_do_epub(blocos, caminho):
    exportar.para_epub([PaginaExtraida(numero=0, blocos=blocos)], caminho,
                       titulo="Teste")
    with zipfile.ZipFile(caminho) as z:
        alvo = [n for n in z.namelist() if n.startswith("OEBPS/pagina")][0]
        return z.read(alvo).decode("utf-8")


def test_a_leitura_da_celula_separa_as_linhas():
    img = np.full((900, 900), 245, np.uint8)
    for k in range(4):
        cv2.putText(img, "ABCDEFGH", (40, 60 + k * 50),
                    cv2.FONT_HERSHEY_SIMPLEX, 1.0, 15, 2)
    cv2.rectangle(img, (60, 300), (840, 700), 15, 3)
    cv2.line(img, (450, 300), (450, 700), 15, 3)
    cv2.line(img, (60, 500), (840, 500), 15, 3)
    for x in (95, 485):
        for y0 in (300, 500):
            cv2.putText(img, "HHH", (x, y0 + 70), cv2.FONT_HERSHEY_SIMPLEX, 1.3, 15, 2)
            cv2.putText(img, "HHH", (x, y0 + 155), cv2.FONT_HERSHEY_SIMPLEX, 0.5, 15, 2)

    def por_altura(recorte):
        if recorte.size == 0:
            return "", 0.0
        return ("A" if recorte.shape[0] > 22 else "b"), 1.0

    boxes = BoxService.generate_boxes_opencv(Image.fromarray(img))
    achado = livro._tabela_da_pagina(img, boxes, por_altura, 0.0, None, 0)
    assert achado is not None
    celulas = [c for fila in achado[0].linhas for c in fila if c]
    assert celulas
    for celula in celulas:
        linhas = celula.split("\n")
        assert len(linhas) == 2, f"a célula não guardou as duas linhas: {celula!r}"
        assert set(linhas[0].replace(" ", "")) == {"A"}
        assert set(linhas[1].replace(" ", "")) == {"b"}


def test_o_epub_escreve_a_quebra_como_br(tmp_path):
    xhtml = _xhtml_do_epub([Tabela([[CELULA, "x"]])], str(tmp_path / "t.epub"))
    assert "Win (1 " in xhtml
    celula = xhtml.split("<td>")[1].split("</td>")[0]
    assert celula.count("<br/>") == 1
    antes, depois = celula.split("<br/>")
    assert "Win" in antes and "Draw" in depois
    assert "\n" not in celula


def test_a_figurina_das_duas_linhas_ganha_a_fonte(tmp_path):
    xhtml = _xhtml_do_epub([Paragrafo("prosa"), Tabela([[CELULA, "x"]])],
                           str(tmp_path / "t.epub"))
    celula = xhtml.split("<td>")[1].split("</td>")[0]
    assert celula.count('<span class="sim">♖</span>') == 2


def test_o_docx_escreve_a_quebra_no_run(tmp_path):
    pytest.importorskip("docx")
    caminho = str(tmp_path / "t.docx")
    exportar.para_docx([PaginaExtraida(numero=0, blocos=[
        Paragrafo("prosa"), Tabela([[CELULA, "x"]])])], caminho, titulo="Teste")
    from docx import Document
    doc = Document(caminho)
    assert doc.tables[0].cell(0, 0).text == CELULA
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 1, \
        "a quebra é de linha, e não um parágrafo a mais"


def test_o_texto_corrido_da_pagina_continua_com_espaco():
    pagina = PaginaExtraida(numero=0, blocos=[Tabela([[CELULA, "x"]])])
    assert pagina.texto == "W: Win (1 ♖e1!) B: Draw (1...♖a2!) x"


def _pagina_com_a_celula():
    return PaginaExtraida(numero=0, blocos=[Tabela([[CELULA, "x"]])])


def test_o_ir_guarda_a_quebra_e_o_html_do_ir_a_escreve():
    from core.editorial_adapters import (documento_para_paginas_extraidas,
                                         paginas_extraidas_para_documento)
    from core.editorial_export import ExportOptions, _block_html, texto_do_bloco

    documento = paginas_extraidas_para_documento([_pagina_com_a_celula()],
                                                 document_id="book")
    bloco = [b for b in documento.pages[0].blocks if b.kind == "table"][0]
    assert bloco.decision.value["rows"][0][0] == CELULA
    # A camada que a busca lê é texto corrido, como era.
    assert "\n" not in texto_do_bloco(bloco)
    assert "Win (1 ♖e1!)<br/>B: Draw" in _block_html(bloco, ExportOptions())
    # E a volta devolve a célula com a quebra.
    de_volta = documento_para_paginas_extraidas(documento)[0]
    assert de_volta.blocos[0].linhas[0][0] == CELULA


def test_o_editor_recebe_a_quebra_como_quebra_antes():
    from core.editor.importar_ir import de_paginas
    from core.editor.modelo import Tabela as TabelaDoEditor

    livro_do_editor, _relatorio = de_paginas([_pagina_com_a_celula()])
    tabelas = [b for c in livro_do_editor.capitulos for b in c.blocos
               if isinstance(b, TabelaDoEditor)]
    assert tabelas
    trechos = tabelas[0].filas[0][0].blocos[0].trechos
    assert [t.texto for t in trechos] == ["W: Win (1 ♖e1!)", "B: Draw (1...♖a2!)"]
    assert [t.quebra_antes for t in trechos] == [False, True]


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
