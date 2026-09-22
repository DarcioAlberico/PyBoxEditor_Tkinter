"""
O lado a jogar deixa de ser convenção calada (item 3 da revisão de 2026-09-18).

São quatro coisas, e os testes seguem essa ordem: a peneira que lê "White to
play" de um texto; os dois caminhos de leitura (a Fase 4 e o `core/livro.py`)
pondo o lado no FEN quando a página o diz; o carimbo no `alt`, na legenda e no
diálogo legado quando ela não diz; e a figura que não pode faltar no arquivo
exportado.

Os dois critérios de aceite escritos no roadmap estão em
`test_legenda_black_to_move_termina_o_fen_em_b` e
`test_nenhum_diagrama_sai_sem_imagem`.
"""

from __future__ import annotations

import numpy as np
import pytest

from core import lado_a_jogar
from core.editorial_export import (EditorialExporter, ExportOptions,
                                   avisos_dos_diagramas, imagem_do_diagrama,
                                   legenda_do_diagrama, revisao_pendente)
from core.editorial_model import (Decision, EditorialBlock, EditorialDocument,
                                  EditorialPage, Evidence, SourceRef)
from core.editorial_pipeline import PageEvidence, ProcessOptions
from core.ocr_phase4 import (BoardCandidate, DiagramProcessor, Phase4Processor,
                             SquareCandidate, SquareRecognizer,
                             legendas_do_diagrama, resolve_position)
from core.ocr_result import LineResult, PageResult
from core.ocr_runtime import CancellationToken


# ----------------------------------------------------------------------
# A peneira
# ----------------------------------------------------------------------

@pytest.mark.parametrize("texto, esperado", [
    ("White to play", "w"),
    ("Black to move", "b"),
    ("Black to play and win", "b"),
    ("White is to move", "w"),
    ("White on the move", "w"),
    ("As brancas jogam", "w"),
    ("pretas a jogar", "b"),
    ("Jogam as negras", "b"),
    ("vez das pretas", "b"),
    ("Juegan las blancas", "w"),
    ("negras juegan", "b"),
    ("♔ to play", "w"),          # a figurina é a cor que ela é
    ("♚ to move", "b"),
])
def test_a_legenda_que_diz_o_lado(texto, esperado):
    lido = lado_a_jogar.ler(texto)
    assert lido.lado == esperado
    assert lido.origem == "legenda"
    assert bool(lido) is True


@pytest.mark.parametrize("texto", [
    "", "Diagram 5-1", "437", "399a =/=",
    "after White’s 20th move",   # prosa que fala de um lance, não da vez
    "A.Yusupov – A.Reuss",
])
def test_a_legenda_que_nao_diz_nada(texto):
    lido = lado_a_jogar.ler(texto)
    assert lido.lado is None and lido.origem == "ausente"
    assert bool(lido) is False


def test_a_chave_de_tabela_com_os_dois_lados_nao_decide():
    """A página 237 do Nunn abre com a chave da tabela, e ela não é de diagrama."""
    lido = lado_a_jogar.ler(
        "In all cases W♖g1, ♙c3 and B♔h7: "
        "W=White to play B=Black to play")
    assert lido.lado is None
    assert lido.origem == "ambigua"


def test_dois_textos_em_volta_lidos_em_separado():
    """
    O cabeçalho e a legenda são lidos um a um: emendados, o "as brancas" que
    termina um e o "jogam" que abre o outro inventariam uma frase.
    """
    assert lado_a_jogar.ler_varios("Ex. 22-4", "Black to move").lado == "b"
    assert lado_a_jogar.ler_varios("as brancas", "jogam bem").lado is None
    assert lado_a_jogar.ler_varios("White to play", "Black to play").origem == "ambigua"


def test_com_lado_troca_o_campo_sem_mexer_no_resto():
    assert (lado_a_jogar.com_lado("8/8/8/8/8/8/8/4K2k w KQ e3 0 1", "b")
            == "8/8/8/8/8/8/8/4K2k b KQ e3 0 1")
    assert lado_a_jogar.com_lado("8/8/8/8/8/8/8/4K2k", "b").endswith(" b - - 0 1")
    assert lado_a_jogar.com_lado("8/8/8/8/8/8/8/4K2k w - - 0 1", None).endswith(" w - - 0 1")
    assert lado_a_jogar.do_fen("8/8/8/8/8/8/8/4K2k b - - 0 1") == "b"
    assert lado_a_jogar.do_fen("8/8/8/8/8/8/8/4K2k") is None


# ----------------------------------------------------------------------
# A Fase 4
# ----------------------------------------------------------------------

def _candidatos(simbolo: str, confianca: float = .95):
    return (SquareCandidate(simbolo, confianca, "test"),
            SquareCandidate(".", 1.0 - confianca, "test"))


class _Reconhecedor(SquareRecognizer):
    name = "fake-squares"

    def __init__(self, values):
        self.values = values

    def recognize(self, image, board, square, context):
        return self.values.get(square, _candidatos("."))


def _tabuleiro():
    return BoardCandidate(
        id="board-0", bbox=(10, 20, 170, 180),
        quad=((10, 20), (170, 20), (170, 180), (10, 180)),
        confidence=.96, source="test", orientation="branca")


def _evidencia():
    return PageEvidence(
        document_id="book", page_index=0, raster=np.zeros((220, 200), dtype=np.uint8),
        text_layer="Texto da página", raster_hash="hash", raster_dpi=300,
        source_kind="memory", text_blocks=[], metadata={})


def _posicao(**kwargs):
    """Um tabuleiro **inteiro**, para `review_required` falar só do que se testa.

    Faltando candidato de casa, a revisão fica obrigatória por isso, e o teste
    do lado a jogar não teria como distinguir uma causa da outra.
    """
    casas = {f"{coluna}{fila}": _candidatos(".")
             for fila in "12345678" for coluna in "abcdefgh"}
    casas["e1"] = _candidatos("K")
    casas["e8"] = _candidatos("k")
    return resolve_position(casas, orientation="branca", **kwargs)


def test_legenda_black_to_move_termina_o_fen_em_b():
    """O critério de aceite do item 3: a legenda decide o campo do FEN."""
    resultado = _posicao(annotations={"legend": "Black to move"})
    assert resultado.fen.endswith(" b - - 0 1")
    assert resultado.side_to_move == "b"
    assert resultado.side_to_move_source == "legend"
    assert "side_to_move_legend" in resultado.reason_codes
    assert lado_a_jogar.AVISO_ASSUMIDO not in resultado.warnings


def test_sem_legenda_o_w_e_convencao_e_sai_declarado():
    resultado = _posicao(annotations={"legend": "Diagram 5-1"})
    assert resultado.fen.endswith(" w - - 0 1")
    assert resultado.side_to_move_source == "assumed"
    assert "side_to_move_assumed" in resultado.reason_codes
    assert lado_a_jogar.AVISO_ASSUMIDO in resultado.warnings
    # Convenção declarada não é pendência: a fila de revisão não ganha 300
    # diagramas por um campo que a página não imprime.
    assert resultado.review_required is False


def test_o_lado_informado_por_quem_chama_manda():
    resultado = _posicao(side_to_move="b", annotations={"legend": "White to play"})
    assert resultado.side_to_move == "b"
    assert resultado.side_to_move_source == "explicit"
    with pytest.raises(ValueError):
        _posicao(side_to_move="x")


def test_a_revisao_informa_o_lado_e_o_aviso_sai():
    resultado = _posicao()
    assert lado_a_jogar.AVISO_ASSUMIDO in resultado.warnings
    revisado = resultado.review({}, side_to_move="b")
    assert revisado.fen.endswith(" b - - 0 1")
    assert revisado.side_to_move_source == "manual"
    assert lado_a_jogar.AVISO_ASSUMIDO not in revisado.warnings
    assert revisado.original_fen == resultado.fen


def test_a_linha_encostada_no_tabuleiro_e_a_legenda_dele():
    page = PageResult("page-0000", lines=[
        LineResult("l1", "White to play", .9, bbox=(12, 185, 160, 200)),
        LineResult("l2", "Um parágrafo da coluna ao lado", .9, bbox=(300, 80, 420, 95)),
        LineResult("l3", "Longe demais para ser legenda", .9, bbox=(12, 400, 160, 415)),
    ])
    assert legendas_do_diagrama((10, 20, 170, 180), page) == ("White to play",)


def test_a_legenda_da_pagina_chega_ao_diagrama_da_fase_4():
    processador = Phase4Processor(
        text_processor=lambda evidence, options, token: PageResult(
            "page-0000", text="Black to move", confidence=.9,
            lines=[LineResult("l1", "Black to move", .9, bbox=(12, 185, 160, 200))],
            metadata={"engine": "fake"}),
        diagram_processor=DiagramProcessor(
            detector=lambda image, **kwargs: [_tabuleiro()],
            recognizer=_Reconhecedor({"e1": _candidatos("K"), "e8": _candidatos("k")})),
    )
    page = processador.process(_evidencia(), ProcessOptions(use_cache=False),
                               CancellationToken())
    diagrama = page.metadata["diagrams"][0]
    assert diagrama["fen"].endswith(" b - - 0 1")
    assert diagrama["side_to_move_source"] == "legend"


# ----------------------------------------------------------------------
# O carimbo e a figura no arquivo exportado
# ----------------------------------------------------------------------

def _documento(valor, *, status="automatic", metadata=None) -> EditorialDocument:
    ref = SourceRef("book", 0, (10, 20, 180, 100), "hash", "raster")
    evidencia = Evidence("ev", ref, observed_text="", confidence=.9)
    bloco = EditorialBlock("d", "diagram", 0, [ref],
                           Decision(valor, ["ev"], status),
                           metadata=dict(metadata or {}))
    return EditorialDocument("book", "Livro", "pt", [
        EditorialPage("book-p0001", 0, [ref], [bloco], [evidencia])],
        pipeline_version="v4")


def _bloco(documento) -> EditorialBlock:
    return documento.pages[0].blocks[0]


def test_nenhum_diagrama_sai_sem_imagem(tmp_path):
    """
    O outro critério de aceite: a Fase 4 não traz pixels — guarda a posição e o
    hash do recorte —, e o `<figure>` dela saía com legenda e sem figura.
    """
    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1",
                            "review_status": "review_required"})
    imagem, origem = imagem_do_diagrama(_bloco(documento))
    assert imagem and origem == "desenho"

    destino = tmp_path / "livro.html"
    relatorio = EditorialExporter().export(
        documento, destino, ExportOptions(format="html", mode="clean"))
    html = destino.read_text(encoding="utf-8")
    assert '<img alt=' in html and "data:image/png;base64," in html
    assert 'data-image="desenho"' in html
    assert not [aviso for aviso in relatorio.warnings if "sem imagem" in aviso]


def test_o_recorte_da_pagina_tem_precedencia_sobre_o_desenho():
    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1",
                            "png_base64": "aGVsbG8="})
    assert imagem_do_diagrama(_bloco(documento)) == ("aGVsbG8=", "recorte")


def test_o_diagrama_sem_fen_nem_imagem_vira_aviso_do_relatorio(tmp_path):
    documento = _documento({"fen": ""})
    assert imagem_do_diagrama(_bloco(documento)) == ("", "nenhuma")
    relatorio = EditorialExporter().export(
        documento, tmp_path / "livro.html", ExportOptions(format="html"))
    assert any("sem imagem" in aviso for aviso in relatorio.warnings)


def test_o_nao_revisado_e_a_convencao_saem_ate_no_modo_limpo(tmp_path):
    """
    O modo limpo tira a proveniência, que é para quem revisa. A ressalva do
    lado a jogar e o "não revisado" são outra coisa: são o que impede uma
    leitura de 94,5% de casas certas de virar afirmação no livro entregue.
    """
    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1",
                            "review_status": "review_required"})
    bloco = _bloco(documento)
    assert revisao_pendente(bloco) is True
    assert legenda_do_diagrama(bloco) == ("brancas a jogar (assumido: a página "
                                          "não diz); não revisado")

    destino = tmp_path / "livro.html"
    relatorio = EditorialExporter().export(
        documento, destino, ExportOptions(format="html", mode="clean"))
    html = destino.read_text(encoding="utf-8")
    assert "Proveniência e auditoria" not in html, "o modo limpo trouxe auditoria"
    assert 'data-review="pending"' in html
    assert "não revisado" in html
    assert "assumido" in html
    assert relatorio.warnings == ("1 diagrama(s) exportados sem revisão",)


def test_o_lado_lido_da_legenda_nao_e_chamado_de_assumido():
    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k b - - 0 1",
                            "side_to_move_source": "legend",
                            "review_status": "reviewed"}, status="reviewed")
    assert legenda_do_diagrama(_bloco(documento)) == "pretas a jogar (da legenda)"
    assert avisos_dos_diagramas(documento) == ()


def test_o_epub_leva_a_figura_do_diagrama(tmp_path):
    import zipfile

    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1"})
    destino = tmp_path / "livro.epub"
    EditorialExporter().export(documento, destino, ExportOptions(format="epub"))
    with zipfile.ZipFile(destino) as arquivo:
        xhtml = arquivo.read("OEBPS/content.xhtml").decode("utf-8")
    assert "data:image/png;base64," in xhtml


def test_o_docx_leva_a_figura_e_a_ressalva(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    documento = _documento({"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1",
                            "review_status": "review_required"})
    destino = tmp_path / "livro.docx"
    EditorialExporter().export(documento, destino, ExportOptions(format="docx"))
    word = Document(destino)
    texto = "\n".join(par.text for par in word.paragraphs)
    assert "não revisado" in texto and "assumido" in texto
    assert word.inline_shapes, "o DOCX saiu sem o tabuleiro"


# ----------------------------------------------------------------------
# O caminho legado: a leitura, o tabuleiro em edição e o diálogo
# ----------------------------------------------------------------------

def _leitura(lado=None):
    from core import diagrama_modelo

    casas = [diagrama_modelo.Casa(0, 4, "k", 1.0),
             diagrama_modelo.Casa(7, 4, "K", 1.0)]
    leitura = diagrama_modelo.Leitura(caixa=(0, 0, 80, 80), casas=casas)
    leitura.lado_a_jogar = lado
    return leitura


def test_a_leitura_sem_legenda_continua_na_convencao():
    assert _leitura().fen().endswith(" w - - 0 1")


def test_a_leitura_com_o_lado_da_legenda_sai_com_ele():
    assert _leitura("b").fen().endswith(" b - - 0 1")


def test_o_titulo_do_diagrama_troca_o_aviso_da_convencao():
    from core import diagrama

    leitura = _leitura()
    leitura.avisos.append(diagrama.AVISO_CONVENCAO)
    leitura.titulo.texto = "Black to play"
    diagrama._lado_do_titulo(leitura)

    assert leitura.lado_a_jogar == "b"
    assert diagrama.AVISO_CONVENCAO not in leitura.avisos
    assert any("vem da legenda" in aviso for aviso in leitura.avisos)


def test_o_tabuleiro_em_edicao_distingue_legenda_de_convencao():
    from core.tabuleiro_edicao import AVISO_CONVENCAO, TabuleiroEdicao

    assumido = TabuleiroEdicao(_leitura())
    assert assumido.lado == "w" and assumido.lado_origem == "convencao"
    assert AVISO_CONVENCAO in assumido.avisos()

    lido = TabuleiroEdicao(_leitura("b"))
    assert lido.lado == "b" and lido.lado_origem == "legenda"
    assert lido.fen().split()[1] == "b"
    assert any("da legenda da página" in aviso for aviso in lido.avisos())
    assert not any("quem editou" in aviso for aviso in lido.avisos())

    # Mexido à mão, o lado passa a ser de quem revisa — e o aviso muda junto.
    lido.definir_lado("w")
    assert lido.lado_origem == "usuario"
    assert any("quem editou" in aviso for aviso in lido.avisos())
    lido.desfazer()
    assert lido.lado == "b" and lido.lado_origem == "legenda"


@pytest.mark.gui
def test_o_dialogo_legado_avisa_a_orientacao_nao_confirmada():
    """
    Sem coordenadas em volta, a leitura assume brancas embaixo — e um diagrama
    impresso do lado das pretas sai plausível, legal e espelhado. Quem confere
    casa a casa não desconfia: a leitura bate com a tela, e as duas estão
    erradas do mesmo jeito.
    """
    from PIL import Image
    from conftest import raiz_tk
    from core import diagrama_modelo
    from ui.dialogo_diagrama import DialogoDiagrama

    raiz = raiz_tk()
    if raiz is None:
        pytest.skip("sem display")
    try:
        imagem = Image.fromarray(np.full((80, 80), 250, np.uint8))
        sem_rotulos = _leitura()
        com_rotulos = _leitura()
        com_rotulos.rotulos = diagrama_modelo.Rotulos(
            lados=("abaixo",), colunas="abcdefgh", filas="87654321",
            orientacao="branca")

        dlg = DialogoDiagrama(raiz, imagem, [sem_rotulos, com_rotulos])
        dlg.construir()
        assert "Orientação não confirmada" in dlg.lbl_avisos.cget("text")
        assert "não há coordenadas impressas" in dlg.lbl_avisos.cget("text")

        dlg._ir(1)
        assert "Orientação não confirmada" not in dlg.lbl_avisos.cget("text")
    finally:
        try:
            raiz.destroy()
        except Exception:
            pass


def test_a_figura_do_livro_leva_o_lado_da_legenda():
    """
    O recorte é o caminho sem modelo neural, e serve para o que se testa aqui:
    a legenda impressa embaixo do tabuleiro decide o campo do FEN e a
    procedência que o `alt` vai declarar.
    """
    from core import diagrama_modelo, livro

    imagem = np.full((200, 200), 250, np.uint8)
    d = livro.Diagrama(exclusao=(10, 10, 190, 190), tabuleiro=(20, 20, 180, 180))
    d.legenda = diagrama_modelo.Titulo(texto="Black to play", lado="abaixo")

    figura = livro._figura_do_diagrama(
        imagem, d, dpi=300, dpi_figura=300, modo="recorte", coordenadas=False,
        fonte="SkakNew-Diagram", lado=192)
    assert figura.lado_a_jogar == "b" and figura.lado_origem == "legenda"

    sem_legenda = livro.Diagrama(exclusao=(10, 10, 190, 190),
                                 tabuleiro=(20, 20, 180, 180))
    figura = livro._figura_do_diagrama(
        imagem, sem_legenda, dpi=300, dpi_figura=300, modo="recorte",
        coordenadas=False, fonte="SkakNew-Diagram", lado=192)
    assert figura.lado_a_jogar is None and figura.lado_origem == "convencao"
