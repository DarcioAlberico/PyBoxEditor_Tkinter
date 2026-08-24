"""
Testes da F99 — a coordenada desenhada pela própria fonte de xadrez.

A Chess Merida tem dezesseis glifos que trazem o filete da moldura **com o
rótulo da fila ou da coluna ao lado**. Com eles o diagrama com coordenada passa
a caber inteiro em texto, e é isso que a fase compra:

    no EPUB    some o `<i>` dentro do `<span>`, que existia só para pôr um
               rótulo de fonte de texto em cima de uma casa da fonte de xadrez
    no DOCX    diagrama com coordenada deixa de cair para imagem — era uma
               limitação escrita no docstring da `para_docx` desde a F59
    no PNG     o rótulo sai no tipo do livro, e não na Helvetica do PyMuPDF

**E a fonte é quem decide.** A SkakNew-Diagram não tem glifo de borda nenhum, e
tudo o que ela faz continua igual: filete de caneta, rótulo em fonte de texto,
imagem no DOCX. Metade destes testes é sobre isso — que o caminho novo não
encoste no caminho velho.

Rodar sem pytest:      python tests/test_f99_coordenada_em_glifo.py
"""

import json
import os
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from core import exportar, livro
from core import render_diagrama as rd

MERIDA = "ChessMerida-Diagram"
SKAK = "SkakNew-Diagram"
FEN = "3qkb2/5p2/2n5/1B2P3/3P1r2/2N5/5P2/2RQK3"


def _figura(nome=MERIDA, moldura="dupla", coordenadas=True, lado=352):
    """Uma figura como o `livro` a monta, com o mesmo critério dele."""
    png, largura, altura = rd.desenhar(FEN, fonte=nome, lado_px=lado,
                                       moldura=moldura,
                                       coordenadas=coordenadas)
    fonte = rd.carregar(nome)
    em_grade = (rd.grade(FEN, fonte, "branca", moldura) if coordenadas
                else None)
    return livro.Figura(png, largura, altura, fen=FEN, origem="render",
                        linhas=(em_grade or rd.linhas(FEN, fonte)),
                        linhas_emolduradas=em_grade is not None,
                        fonte=nome, coordenadas=coordenadas,
                        casas_de_largura=largura * 8.0 / rd.lado_efetivo(lado))


def _pagina(*figuras):
    return [livro.PaginaExtraida(numero=0, blocos=list(figuras))]


# ----------------------------------------------------------------------
# A grade
# ----------------------------------------------------------------------

def test_a_grade_e_dez_por_dez_e_o_miolo_e_o_tabuleiro():
    fonte = rd.carregar(MERIDA)
    grade = rd.grade(FEN, fonte, "branca", "dupla")
    assert len(grade) == 10
    assert all(len(linha) == 10 for linha in grade)
    # O miolo, sem a coluna da esquerda e a da direita, é o tabuleiro nu.
    assert [linha[1:-1] for linha in grade[1:-1]] == rd.linhas(FEN, fonte)


def test_a_fonte_sem_glifo_de_borda_devolve_nada():
    """
    **Devolver `None` é a resposta, e não uma falha.** A SkakNew-Diagram tem 46
    codepoints e nenhum deles é filete; quem a usa continua no caminho de
    sempre.
    """
    assert rd.grade(FEN, rd.carregar(SKAK), "branca", "dupla") is None


def test_sem_moldura_nao_ha_grade():
    """
    Na Merida o rótulo vem **dentro** do glifo do filete: não há um sem o
    outro. Pedir coordenada sem moldura continua saindo pela caneta.
    """
    assert rd.grade(FEN, rd.carregar(MERIDA), "branca", "sem") is None


def test_a_grade_gira_com_o_tabuleiro():
    """
    O diagrama impresso do lado das pretas tem `h` na primeira coluna e `1` na
    primeira fila. O rótulo sai dos mesmos `rotulos` que o resto usa — sem isso,
    o `a1` do desenho sairia rotulado `h8` e nada denunciaria.
    """
    fonte = rd.carregar(MERIDA)
    pecas = fonte.moldura_em_glifo("dupla")
    grade = rd.grade(FEN, fonte, "preta", "dupla")

    assert grade[1][0] == pecas["filas"][0], "a fila de cima devia ser a 1"
    assert grade[-1][1] == pecas["colunas"][7], "a coluna da esquerda devia ser h"

    de_pe = rd.grade(FEN, fonte, "branca", "dupla")
    assert de_pe[1][0] == pecas["filas"][7]
    assert de_pe[-1][1] == pecas["colunas"][0]


def test_o_glifo_de_borda_entra_na_conferencia_de_cobertura():
    """
    Mesma disciplina das casas (SPEC §4.2): um mapa que prometa um filete que a
    fonte não desenha daria a coluna dos rótulos em branco — plausível à
    distância. A cobaia é a SkakNew, que não tem `0xC0`.
    """
    mapas = json.load(open(rd.CAMINHO_MAPAS, encoding="utf-8"))
    mapas["fontes"]["Cobaia"] = dict(
        mapas["fontes"][SKAK],
        molduras={"dupla": dict(mapas["fontes"][MERIDA]["molduras"]["dupla"])})
    with tempfile.TemporaryDirectory() as tmp:
        caminho = os.path.join(tmp, "mapas.json")
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(mapas, f)
        original = rd.CAMINHO_MAPAS
        rd.CAMINHO_MAPAS = caminho
        rd.esquecer()
        try:
            with pytest.raises(rd.FonteIncompleta):
                rd.carregar("Cobaia")
        finally:
            rd.CAMINHO_MAPAS = original
            rd.esquecer()


# ----------------------------------------------------------------------
# O desenho
# ----------------------------------------------------------------------

def test_o_png_com_coordenada_sai_recortado_na_tinta():
    """
    A grade tem dez casas de lado, mas a tinta não chega às bordas dela — o
    filete de cima mora no pé da casa de cima. Sem recortar, o desenho sairia
    com quase uma casa de branco em cima e à direita e nada embaixo.
    """
    _png, largura, altura = rd.desenhar(FEN, fonte=MERIDA, lado_px=352,
                                        moldura="dupla", coordenadas=True)
    casa = 44.0
    assert 8.5 < largura / casa < 9.6, f"{largura / casa:.2f} casas de largura"
    assert 8.5 < altura / casa < 9.6

    # O recorte é determinístico: a moldura fecha o desenho dos quatro lados e
    # as oito filas e colunas saem sempre, então a caixa da tinta não depende de
    # onde estão as peças.
    outra = rd.desenhar("8/8/8/8/8/8/8/8", fonte=MERIDA, lado_px=352,
                        moldura="dupla", coordenadas=True)
    assert outra[1:] == (largura, altura)


def test_sem_coordenada_nada_muda_em_nenhuma_das_duas_fontes():
    """O caminho novo não pode encostar no velho — é metade desta fase."""
    for moldura in ("sem", "simples", "dupla"):
        lados = {rd.desenhar(FEN, fonte=n, lado_px=352, moldura=moldura)[1]
                 for n in rd.fontes()}
        assert len(lados) == 1, f"as fontes divergiram em {moldura}: {lados}"


def test_a_skaknew_com_coordenada_continua_na_caneta():
    """Filete desenhado e rótulo em fonte de texto, como desde a F58."""
    com = rd.desenhar(FEN, fonte=SKAK, lado_px=352, moldura="simples",
                      coordenadas=True)[1]
    sem = rd.desenhar(FEN, fonte=SKAK, lado_px=352, moldura="simples")[1]
    assert com > sem, "a calha do rótulo sumiu"


# ----------------------------------------------------------------------
# O arquivo escrito
# ----------------------------------------------------------------------

def test_no_epub_o_diagrama_emoldurado_dispensa_o_span_e_a_borda():
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_epub(_pagina(_figura()), os.path.join(tmp, "a.epub"),
                                 diagramas="fonte", moldura="dupla")
    with zipfile.ZipFile(caminho) as z:
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

    assert "<img" not in pagina, "o diagrama com coordenada saiu como imagem"
    assert 'class="rot"' not in pagina and 'class="col"' not in pagina
    # Sem a classe `caixa`: a moldura já está dentro do texto, e a borda da CSS
    # por cima seria a segunda.
    assert re.search(r'<div class="diagrama fonte-[\w-]+"', pagina)
    assert pagina.count("<p>") == 10


def test_no_docx_o_diagrama_com_coordenada_deixa_de_virar_imagem():
    """
    Era limitação escrita no docstring da `para_docx` desde a F59, e o motivo
    dela era a fonte, não o formato.
    """
    from docx import Document

    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(_pagina(_figura()), os.path.join(tmp, "a.docx"),
                                 diagramas="fonte", moldura="dupla",
                                 corpo_pt=16)
    doc = Document(caminho)
    assert not doc.inline_shapes, "continuou saindo como imagem"
    assert len(doc.tables) == 1
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 10

    with zipfile.ZipFile(caminho) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    # Dez casas de 16 pt são 160 pt, que são 3200 twips.
    assert '<w:tblW w:type="dxa" w:w="3200"/>' in xml
    # E a caixa não põe moldura: o texto já traz a sua.
    assert 'w:val="nil"' in xml and "double" not in xml


def test_no_docx_a_skaknew_com_coordenada_continua_em_imagem():
    from docx import Document

    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(_pagina(_figura(SKAK)),
                                 os.path.join(tmp, "b.docx"),
                                 diagramas="fonte", moldura="dupla")
    doc = Document(caminho)
    assert len(doc.inline_shapes) == 1
    assert not doc.tables


def test_duas_fontes_no_mesmo_livro_nao_se_atropelam_na_css():
    """
    **Enquanto havia uma fonte só isto não podia falhar, e agora pode.** A
    família morava em `div.diagrama p`, repetida uma vez por fonte: com duas, a
    segunda cópia venceria em cascata e o livro inteiro sairia na última fonte
    declarada — inclusive os diagramas desenhados com a outra.
    """
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_epub(
        _pagina(_figura(MERIDA, coordenadas=False), _figura(SKAK, coordenadas=False)),
        os.path.join(tmp, "c.epub"), diagramas="fonte")
    with zipfile.ZipFile(caminho) as z:
        css = z.read("OEBPS/estilo.css").decode("utf-8")
        pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

    assert css.count("div.diagrama p {") == 1, "a regra geral saiu duplicada"
    for nome in (MERIDA, SKAK):
        classe = exportar.classe_da_fonte(nome)
        assert f'div.diagrama.{classe} p {{ font-family: "{nome}"' in css
        assert classe in pagina


def test_a_classe_da_fonte_nao_deixa_passar_o_que_a_css_nao_aceita():
    assert exportar.classe_da_fonte("SkakNew-Diagram") == "fonte-SkakNew-Diagram"
    assert exportar.classe_da_fonte("Chess Merida!") == "fonte-Chess-Merida-"


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
