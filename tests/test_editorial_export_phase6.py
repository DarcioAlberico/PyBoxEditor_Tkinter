from __future__ import annotations

import zipfile

import fitz
import pytest

from core.editorial_model import Decision, EditorialBlock, EditorialDocument, EditorialPage, Evidence, SourceRef
from core.editorial_export import EditorialExporter, ExportOptions


def _document(tmp_path=None) -> EditorialDocument:
    ref = SourceRef("book", 0, (10, 20, 180, 100), "hash", "raster")
    evidence = Evidence("ev", ref, observed_text="Livro", confidence=.9)
    blocks = [
        EditorialBlock("h", "heading", 0, [ref], Decision("Capítulo 1", ["ev"], "reviewed")),
        EditorialBlock("p", "paragraph", 1, [ref], Decision("Texto editorial", ["ev"], "reviewed")),
        EditorialBlock("n", "chess_sequence", 2, [ref], Decision("1. e4 e5", ["ev"], "reviewed")),
        EditorialBlock("d", "diagram", 3, [ref], Decision(
            {"fen": "8/8/8/8/8/8/8/4K2k w - - 0 1", "review_status": "reviewed"},
            ["ev"], "reviewed")),
    ]
    return EditorialDocument("book", "Livro de xadrez", "pt", [
        EditorialPage("book-p0001", 0, [ref], blocks, [evidence])
    ], pipeline_version="v4")


def test_html_semantico_preserva_ordem_fen_alt_e_modo(tmp_path):
    destino = tmp_path / "livro.html"
    report = EditorialExporter().export(
        _document(), destino, ExportOptions(format="html", mode="hybrid"))
    html = destino.read_text(encoding="utf-8")

    assert report.format == "html"
    assert '<html lang="pt">' in html
    assert 'data-fen="8/8/8/8/8/8/8/4K2k w - - 0 1"' in html
    assert 'aria-label="Sequência de xadrez"' in html
    assert "Texto editorial" in html
    assert "hybrid" in html


def test_epub3_tem_mimetype_nav_xhtml_css_e_ordem_de_blocos(tmp_path):
    destino = tmp_path / "livro.epub"
    report = EditorialExporter().export(_document(), destino, ExportOptions(format="epub"))

    assert report.files == (str(destino),)
    with zipfile.ZipFile(destino) as arquivo:
        assert arquivo.namelist()[0] == "mimetype"
        assert arquivo.read("mimetype") == b"application/epub+zip"
        xhtml = arquivo.read("OEBPS/content.xhtml").decode("utf-8")
        assert "Capítulo 1" in xhtml and "Texto editorial" in xhtml
        assert "OEBPS/nav.xhtml" in arquivo.namelist()
        assert "OEBPS/styles.css" in arquivo.namelist()


def test_docx_mapeia_heading_notacao_e_diagrama(tmp_path):
    pytest.importorskip("docx")
    destino = tmp_path / "livro.docx"
    EditorialExporter().export(_document(), destino, ExportOptions(format="docx"))
    from docx import Document
    doc = Document(destino)
    texto = "\n".join(par.text for par in doc.paragraphs)
    assert "Capítulo 1" in texto
    assert "1. e4 e5" in texto
    assert "4K2k" in texto


def test_pdf_semantico_e_pesquisavel_tem_relatorio_de_itens(tmp_path):
    destino = tmp_path / "livro.pdf"
    report = EditorialExporter().export(_document(), destino, ExportOptions(format="pdf"))
    assert report.metadata["text_items"] >= 3
    documento = fitz.open(destino)
    texto = documento[0].get_text()
    documento.close()
    assert "Texto editorial" in texto and "1. e4 e5" in texto


def test_modo_invalido_e_formato_invalido_sao_rejeitados():
    with pytest.raises(ValueError):
        ExportOptions(format="epub", mode="inexistente")
    with pytest.raises(ValueError):
        ExportOptions(format="xlsx")
