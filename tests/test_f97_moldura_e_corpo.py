"""
Testes da F97 — a moldura do diagrama e o corpo em pontos.

Duas escolhas que o usuário passa a fazer na hora de exportar, e três lugares
onde elas podem sair diferentes do que ele pediu: o PNG desenhado, o EPUB e o
DOCX. O que estes testes prendem é que as três respondem à **mesma** escolha.

**O ponto mais frágil é o vão entre as filas do DOCX, e ele é o motivo da
fase.** A casa é o quadrado do em da fonte, então o tabuleiro em texto só sai
quadrado se a entrelinha medir exatamente o corpo. O Word escreve o corpo em
meios-pontos (`w:sz`) e a entrelinha em twips (`w:line`), e um corpo que não
caia no meio ponto faz os dois campos dizerem números diferentes — 0,2 pt por
fila, oito filas, e o tabuleiro deixa de fechar. Isso não aparece em nada que
não seja abrir o arquivo no Word, e por isso está preso aqui em XML.

O resto é estrutura: a margem que a moldura ocupa no PNG, a borda na CSS e no
`w:tcBorders`, e a largura da figura, que passou a sair do corpo em vez de uma
constante em centímetros.

Rodar sem pytest:      python tests/test_f97_moldura_e_corpo.py
"""

import os
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from core import exportar, livro
from core import render_diagrama as rd

FEN = "r1bqk2r/pp2bppp/2n1pn2/3p4/3P4/2N1PN2/PP2BPPP/R1BQK2R w - - 0 1"

#: Múltiplo de 8, para o lado do desenho não ser arredondado por baixo do teste.
LADO = 352


def _figura(moldura="simples", texto=True, lado=LADO):
    """Uma figura de diagrama, como o `livro` a monta."""
    png, largura, altura = rd.desenhar(FEN, lado_px=lado, moldura=moldura)
    em_casas = largura * 8.0 / rd.lado_efetivo(lado)
    if not texto:
        # O recorte justo é o retângulo do tabuleiro, e nada mais: são as oito
        # casas redondas, que é o que o `_figura_do_diagrama` calcula para ele.
        return livro.Figura(png, largura, altura, origem="recorte",
                            casas_de_largura=8.0)
    return livro.Figura(png, largura, altura, fen=FEN, origem="render",
                        linhas=rd.linhas(FEN), fonte=rd.FONTE_PADRAO,
                        casas_de_largura=em_casas)


def _paginas(*figuras):
    return [livro.PaginaExtraida(numero=0, blocos=list(figuras))]


def _docx(paginas, **kw):
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_docx(paginas, os.path.join(tmp, "livro.docx"), **kw)
    with zipfile.ZipFile(caminho) as z:
        return caminho, z.read("word/document.xml").decode("utf-8")


def _css(paginas, **kw):
    tmp = tempfile.mkdtemp()
    caminho = exportar.para_epub(paginas, os.path.join(tmp, "livro.epub"), **kw)
    with zipfile.ZipFile(caminho) as z:
        return z.read("OEBPS/estilo.css").decode("utf-8")


# ----------------------------------------------------------------------
# O desenho
# ----------------------------------------------------------------------

def test_a_moldura_cresce_de_sem_para_simples_para_dupla():
    """
    Os três feitios se distinguem pela margem, e o tabuleiro não muda.

    É a prova de que o filete mora **fora** do tabuleiro nos três casos: se ele
    invadisse a casa, o desenho teria o mesmo lado nos três e a diferença
    apareceria dentro do tabuleiro, onde quem relê o diagrama divide a imagem
    em 8×8 iguais.
    """
    lados = [rd.desenhar(FEN, lado_px=LADO, moldura=m)[1]
             for m in ("sem", "simples", "dupla")]
    assert lados[0] == LADO
    assert lados[0] < lados[1] < lados[2]


def test_a_moldura_dupla_desenha_dois_filetes():
    """
    Duas transições de tinta na margem, e não uma.

    Lido na coluna do meio do desenho, de fora para dentro: fundo, filete
    grosso, vão, filete fino, tabuleiro. É o que separa a dupla de uma simples
    mais gorda, e é a única diferença que o usuário pediu para ver.
    """
    from PIL import Image
    import io as _io

    png, largura, _altura = rd.desenhar(FEN, lado_px=LADO, moldura="dupla",
                                        tons=0)
    imagem = Image.open(_io.BytesIO(png)).convert("L")
    coluna = [imagem.getpixel((largura // 2, y)) for y in range(largura // 2)]
    escuro = [p < 128 for p in coluna]
    viradas = sum(1 for a, b in zip(escuro, escuro[1:]) if a != b)
    # fundo→grosso, grosso→vão, vão→fino, fino→tabuleiro: quatro viradas antes
    # de chegar ao meio do tabuleiro (a primeira fila é escura ou clara, e a
    # contagem para na metade da imagem).
    assert viradas >= 4, f"a moldura dupla saiu com {viradas} transições"


def test_os_booleanos_de_antes_continuam_valendo():
    """`True` e `False` são o que a `desenhar` aceitava, e ainda aceita."""
    assert rd.normalizar_moldura(True) == "simples"
    assert rd.normalizar_moldura(False) == "sem"
    assert (rd.desenhar(FEN, lado_px=LADO, moldura=True)[1]
            == rd.desenhar(FEN, lado_px=LADO, moldura="simples")[1])


def test_moldura_escrita_errado_reclama_em_vez_de_virar_simples():
    """Um `"Dupla"` com maiúscula não pode sair como moldura simples calada."""
    for valor in ("Dupla", "nenhuma", 2):
        with pytest.raises(ValueError):
            rd.normalizar_moldura(valor)


# ----------------------------------------------------------------------
# O corpo
# ----------------------------------------------------------------------

def test_o_corpo_cai_no_meio_ponto():
    """
    É o arredondamento que mantém o tabuleiro quadrado no Word.

    Ver `PASSO_DO_CORPO_PT`: `w:sz` é em meios-pontos e `w:line` em twips, e um
    corpo fora do passo faz os dois campos dizerem medidas diferentes.
    """
    assert exportar.corpo_valido(16) == 16.0
    assert exportar.corpo_valido(16.3) == 16.5
    assert exportar.corpo_valido(16.2) == 16.0
    assert exportar.corpo_valido("18") == 18.0


def test_corpo_impossivel_reclama():
    for valor in (0, -3, "grande", None):
        with pytest.raises(ValueError):
            exportar.corpo_valido(valor)


def test_no_docx_a_entrelinha_e_exatamente_o_corpo():
    """
    **O teste da fase.** Sem isto o tabuleiro sai com faixa branca entre as
    filas, e é o defeito que se vê antes de qualquer outro.

    `w:sz` conta meios-pontos e `w:line` conta twips (1/20 pt): para 16 pt são
    32 e 320, e `w:lineRule="exact"` é o que impede o Word de acrescentar o vão
    da fonte por cima.
    """
    _caminho, xml = _docx(_paginas(_figura()), diagramas="fonte", corpo_pt=16)

    espacamento = re.search(r'<w:spacing [^>]*w:line="(\d+)" '
                            r'w:lineRule="(\w+)"', xml)
    assert espacamento, "o parágrafo do tabuleiro saiu sem entrelinha declarada"
    assert espacamento.group(2) == "exact"
    assert int(espacamento.group(1)) == 320

    corpo = re.search(r'<w:sz w:val="(\d+)"/>', xml)
    assert corpo and int(corpo.group(1)) == 32
    assert 'w:before="0"' in xml and 'w:after="0"' in xml


def test_a_figura_sai_no_tamanho_pedido_nos_dois_formatos():
    """
    Oito casas de tabuleiro mais a moldura, e não uma largura fixa.

    O desenho com moldura simples tem 8,08 casas de largura; a 16 pt por casa
    isso são ~129 pt, que é o que tem de chegar ao DOCX e ao EPUB.
    """
    from docx import Document

    figura = _figura("simples")
    esperado = figura.casas_de_largura * 16

    caminho, _xml = _docx(_paginas(figura), corpo_pt=16)
    largura = Document(caminho).inline_shapes[0].width.pt
    assert abs(largura - esperado) < 0.5

    tmp = tempfile.mkdtemp()
    epub = exportar.para_epub(_paginas(figura), os.path.join(tmp, "l.epub"),
                              corpo_pt=16)
    with zipfile.ZipFile(epub) as z:
        xhtml = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")
    achado = re.search(r'style="width:([\d.]+)pt"', xhtml)
    assert achado and abs(float(achado.group(1)) - esperado) < 0.5


def test_o_recorte_sai_do_tamanho_do_desenho():
    """
    Os dois convivem no mesmo livro, e um maior que o outro denunciaria em que
    página o modelo se saiu bem — que é o que a F58 já evitava no rótulo.

    Sem moldura os dois medem oito casas redondas e a igualdade é exata; com
    moldura o desenho leva o filete junto, e a diferença tem de ser **só** ele.
    """
    recorte = _figura(texto=False)
    assert (exportar.largura_em_pt(_figura("sem"), 16)
            == exportar.largura_em_pt(recorte, 16) == 128.0)

    com_filete = exportar.largura_em_pt(_figura("simples"), 16)
    assert 128.0 < com_filete < 128.0 * 1.02


def test_pagina_inteira_em_imagem_nao_ganha_largura_de_tabuleiro():
    """Não há tabuleiro por dentro dela; 8 casas seriam uma miniatura."""
    pagina = livro.Figura(b"", 1000, 1400, origem="pagina")
    assert exportar.largura_em_pt(pagina, 16) is None


# ----------------------------------------------------------------------
# A moldura no arquivo escrito
# ----------------------------------------------------------------------

def test_a_css_do_epub_leva_a_moldura_e_o_corpo_escolhidos():
    css = _css(_paginas(_figura("dupla")), diagramas="fonte", corpo_pt=20,
               moldura="dupla")
    assert "double" in css
    assert "font-size: 20pt" in css
    # A entrelinha do tabuleiro se defende do leitor que impõe a dele: sem isso
    # a fila descola da fila, que é o mesmo defeito do DOCX.
    assert "line-height: 1 !important" in css

    sem = _css(_paginas(_figura("sem")), diagramas="fonte", moldura="sem")
    assert "double" not in sem and "solid #000" not in sem


def test_no_docx_a_moldura_e_uma_celula_em_volta_do_tabuleiro():
    """
    Tabela de uma célula, e não borda de parágrafo: esta emolduraria a coluna
    de texto inteira, e o tabuleiro é estreito e centrado.
    """
    from docx import Document

    caminho, xml = _docx(_paginas(_figura("dupla")), diagramas="fonte",
                         moldura="dupla")
    assert 'w:val="double"' in xml
    doc = Document(caminho)
    assert len(doc.tables) == 1
    # As oito filas, e nem uma linha em branco por cima delas.
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 8

    caminho, xml = _docx(_paginas(_figura("sem")), diagramas="fonte",
                         moldura="sem")
    assert "w:tcBorders" not in xml
    assert not Document(caminho).tables


def test_dois_diagramas_seguidos_nao_caem_na_mesma_moldura():
    """
    A página de exercícios é dois tabuleiros sem prosa entre eles.

    Duas `<w:tbl>` coladas no XML o Word abre como **uma** tabela de duas filas
    — os dois diagramas dentro da mesma moldura, um por cima do outro. O
    parágrafo de 1 pt entre elas é o que separa, e é ele que este teste prende.
    """
    from docx import Document

    pagina = _paginas(_figura("dupla"), _figura("dupla"))
    caminho, xml = _docx(pagina, diagramas="fonte", moldura="dupla")

    assert "</w:tbl><w:tbl>" not in xml, "as duas molduras saíram coladas"
    doc = Document(caminho)
    assert len(doc.tables) == 2
    assert [len(t.rows) for t in doc.tables] == [1, 1]


def test_moldura_errada_reclama_nos_dois_formatos():
    with tempfile.TemporaryDirectory() as tmp:
        for escrever, ext in ((exportar.para_epub, "epub"),
                              (exportar.para_docx, "docx")):
            with pytest.raises(ValueError) as erro:
                escrever(_paginas(_figura()), os.path.join(tmp, f"x.{ext}"),
                         moldura="grossa")
            assert "grossa" in str(erro.value)


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
