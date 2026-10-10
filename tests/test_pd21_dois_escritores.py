"""
Dois escritores de EPUB e DOCX, e não três (PD-21, 2026-10-10).

O escritor do IR (`editorial_export._epub`/`_docx`) era o terceiro: um `content.xhtml`
só, as imagens em base64, sem a fonte dos símbolos nem o diagrama redesenhado — e o
único a carimbar "não revisado" no diagrama que a fila não conferiu (§4.6 da
revisão de 2026-09-18). Agora o carimbo mora na `Figura` (`revisao_pendente`) e sai
pelo escritor de produção, e EPUB e DOCX do documento editorial saem dele, pela
fachada, sobre as páginas do leitor com a revisão aplicada ou de volta do IR.

Os aceites do roadmap: o EPUB do IR de um diagrama `review_required` traz o sinal no
`alt` e na `figcaption` (`test_o_epub_do_ir_carimba_o_nao_revisado_no_alt_e_na_figcaption`,
em `tests/test_lado_a_jogar.py`); `tests/test_aceitacao_formatos.py` passa sem
mudar; e `core/` tem dois escritores de EPUB (`test_o_core_tem_dois_escritores_de_epub`).
"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest

from core import exportar, livro
from core.editorial_adapters import (pagina_editorial_para_extraida,
                                     paginas_extraidas_para_documento)
from core.editorial_export import EditorialExporter, ExportOptions
from core.editorial_legacy import OpcoesDeFigura, _redesenhar, aplicar_revisao
from core.editorial_model import (Decision, EditorialBlock, EditorialDocument,
                                  EditorialPage, Evidence, SourceRef)
from core.editorial_pipeline import EditorialPipeline
from core.editorial_pipeline import ExportOptions as OpcoesDaFachada

RAIZ = Path(__file__).resolve().parents[1]
FEN = "8/1k6/1p6/1K6/P1P5/8/8/8 w - - 0 1"
RESSALVA = "brancas a jogar (assumido: a página não diz); não revisado"


def _figura(**campos) -> livro.Figura:
    """Um diagrama desenhado como `aplicar_revisao` o redesenharia: é o que faz o
    EPUB das páginas do leitor e o das páginas de volta do IR saírem iguais."""
    base = livro.Figura(b"", 0, 0, fen=FEN, origem="render", **campos)
    return _redesenhar(base, FEN, OpcoesDeFigura())


def _pagina(**campos) -> livro.PaginaExtraida:
    return livro.PaginaExtraida(numero=0, blocos=[
        livro.Paragrafo("Chapter 1", titulo=True, nivel=1),
        livro.Paragrafo("White to play and win."),
        _figura(**campos),
        livro.Paragrafo("The king walks."),
    ], largura=1000, altura=1400, dpi=200)


def _xhtml(caminho) -> str:
    with zipfile.ZipFile(caminho) as arquivo:
        return arquivo.read("OEBPS/pagina-0001.xhtml").decode("utf-8")


# ----------------------------------------------------------------------
# O aceite estrutural
# ----------------------------------------------------------------------

def test_o_core_tem_dois_escritores_de_epub():
    """Quem grava o `mimetype` de EPUB em `core/`: o escritor de produção e o do
    editor de livros (ED-12, "convive"). O terceiro não existe mais."""
    escritores = sorted(
        str(caminho.relative_to(RAIZ)).replace(os.sep, "/")
        for caminho in (RAIZ / "core").rglob("*.py")
        if "application/epub+zip" in caminho.read_text(encoding="utf-8"))
    assert escritores == ["core/editor/epub.py", "core/exportar.py"]


# ----------------------------------------------------------------------
# O carimbo no escritor de produção
# ----------------------------------------------------------------------

def test_a_figura_pendente_leva_a_ressalva_no_alt_e_na_figcaption_do_epub(tmp_path):
    exportar.para_epub([_pagina(revisao_pendente=True)], str(tmp_path / "p.epub"),
                       titulo="Livro")
    xhtml = _xhtml(tmp_path / "p.epub")
    assert f'alt="{FEN} — {RESSALVA}"' in xhtml
    assert "<figure><img" in xhtml and f"<figcaption>{RESSALVA}</figcaption></figure>" in xhtml

    # No modo de fonte o `div.diagrama` ganha o mesmo `<figure>` com a legenda, e
    # o `title` que o editor lê de volta leva a ressalva atrás do FEN.
    exportar.para_epub([_pagina(revisao_pendente=True)], str(tmp_path / "f.epub"),
                       titulo="Livro", diagramas="fonte")
    xhtml = _xhtml(tmp_path / "f.epub")
    assert '<figure><div class="diagrama' in xhtml
    assert f"</div><figcaption>{RESSALVA}</figcaption></figure>" in xhtml
    assert f'title="{FEN} — {RESSALVA}"' in xhtml

    # Sem estado (a exportação direta da leitura) e com a posição conferida:
    # nada muda no arquivo, nem o `alt` nem a figura.
    for valor in (None, False):
        caminho = str(tmp_path / f"{valor}.epub")
        exportar.para_epub([_pagina(revisao_pendente=valor)], caminho, titulo="Livro")
        xhtml = _xhtml(caminho)
        assert "não revisado" not in xhtml and "<figcaption>" not in xhtml
        assert f'alt="{FEN} — brancas a jogar (assumido: a página não diz)"' in xhtml
    assert (tmp_path / "None.epub").read_bytes() == (tmp_path / "False.epub").read_bytes()


def test_a_figura_pendente_leva_a_ressalva_na_legenda_do_docx(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    exportar.para_docx([_pagina(revisao_pendente=True)], str(tmp_path / "p.docx"),
                       titulo="Livro")
    doc = Document(tmp_path / "p.docx")
    assert RESSALVA in [p.text for p in doc.paragraphs]
    [forma] = doc.inline_shapes
    assert forma._inline.docPr.get("descr") == f"{FEN} — {RESSALVA}"

    exportar.para_docx([_pagina(revisao_pendente=True)], str(tmp_path / "f.docx"),
                       titulo="Livro", diagramas="fonte")
    doc = Document(tmp_path / "f.docx")
    assert RESSALVA in [p.text for p in doc.paragraphs], "no modo de fonte a legenda some"

    exportar.para_docx([_pagina()], str(tmp_path / "n.docx"), titulo="Livro")
    assert not any("não revisado" in p.text for p in Document(tmp_path / "n.docx").paragraphs)


# ----------------------------------------------------------------------
# De onde o estado vem: a revisão, nas cópias
# ----------------------------------------------------------------------

def test_aplicar_revisao_carimba_o_estado_do_diagrama_na_copia_e_nao_no_leitor():
    paginas = [_pagina()]
    documento = paginas_extraidas_para_documento(paginas, document_id="livro")
    # O diagrama desenhado sai automático, sem nada a acusar: fica como o
    # leitor o deixou, o mesmo objeto, sem estado.
    [copia] = aplicar_revisao(paginas, documento)
    assert copia.blocos[2] is paginas[0].blocos[2]
    assert copia.blocos[2].revisao_pendente is None
    # A fila o marca para revisão: a cópia sabe; a página do leitor, não.
    bloco = next(b for b in documento.pages[0].blocks if b.kind == "diagram")
    bloco.metadata["review_status"] = "review_required"
    [copia] = aplicar_revisao(paginas, documento)
    assert copia.blocos[2].revisao_pendente is True
    assert paginas[0].blocos[2].revisao_pendente is None
    assert copia.blocos[2] is not paginas[0].blocos[2]
    # E de volta do IR, sem as páginas do leitor, o mesmo estado chega.
    [de_volta] = aplicar_revisao([], documento)
    assert de_volta.blocos[2].revisao_pendente is True
    # A volta pura continua sendo o inverso exato da ida.
    assert pagina_editorial_para_extraida(documento.pages[0]).blocos[2].revisao_pendente is None


def test_o_diagrama_da_fase_4_volta_desenhado_do_fen_e_a_figura_vazia_fica_fora():
    """A Fase 4 guarda a posição e o hash do recorte, não os pixels: a volta o
    desenha do FEN, para o escritor de produção ter o que escrever. A figura
    sem imagem nem posição não tem o que devolver."""
    ref = SourceRef("book", 0, (10, 20, 180, 100), "hash", "raster")
    pagina = EditorialPage("book-p0001", 0, [ref], [
        EditorialBlock("d", "diagram", 0, [ref], Decision(
            {"fen": FEN, "review_status": "review_required"}, ["ev"], "automatic")),
        EditorialBlock("f", "figure", 1, [ref], Decision({"warning": "sem recorte"},
                                                         ["ev"], "automatic")),
        EditorialBlock("p", "paragraph", 2, [ref], Decision("Texto", ["ev"], "automatic")),
    ], [Evidence("ev", ref, observed_text="", confidence=.9)])
    volta = pagina_editorial_para_extraida(pagina, estado_da_revisao=True)
    figura, paragrafo = volta.blocos
    assert isinstance(figura, livro.Figura) and paragrafo.texto == "Texto"
    assert figura.png[:8] == b"\x89PNG\r\n\x1a\n"
    assert figura.largura == figura.altura > 0 and figura.casas_de_largura > 8
    assert figura.fen == FEN and figura.origem == "render"
    assert figura.revisao_pendente is True
    assert pagina_editorial_para_extraida(pagina).blocos[0].revisao_pendente is None


# ----------------------------------------------------------------------
# Um caminho só: a fachada, com ou sem leitor
# ----------------------------------------------------------------------

def test_a_fachada_escreve_o_mesmo_epub_das_paginas_do_leitor_e_de_volta_do_ir(
        tmp_path, monkeypatch):
    """
    O que a dobra promete: o EPUB da fachada com o leitor (as páginas lidas, a
    revisão aplicada), o da fachada sem leitor (o IR de outra sessão) e o do
    exportador sozinho são o mesmo arquivo, byte a byte — e o mesmo que
    `exportar.para_epub` escreve direto das páginas.
    """
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "1700000000")
    paginas = [_pagina()]
    documento = paginas_extraidas_para_documento(paginas, document_id="livro",
                                                 title="Livro", language="en")

    class _Extrator:
        ultimas_paginas = paginas

    com_leitor = EditorialPipeline(legacy_extractor=_Extrator()).export(
        documento, tmp_path / "leitor.epub", OpcoesDaFachada(format="epub"))
    sem_leitor = EditorialPipeline().export(
        EditorialDocument.from_dict(documento.to_dict()), tmp_path / "ir.epub",
        OpcoesDaFachada(format="epub"))
    exportador = EditorialExporter().export(
        documento, tmp_path / "exportador.epub", ExportOptions(format="epub"))
    assert [r.metadata["paginas"] for r in (com_leitor, sem_leitor, exportador)] == [
        "leitor", "ir", "ir"]
    assert all(r.metadata["escritor"] == "exportar" for r in (com_leitor, sem_leitor, exportador))

    direto = exportar.para_epub(paginas, str(tmp_path / "direto.epub"), titulo="Livro",
                                idioma="en")
    conteudos = [Path(p).read_bytes() for p in (
        tmp_path / "leitor.epub", tmp_path / "ir.epub", tmp_path / "exportador.epub", direto)]
    assert conteudos[0] == conteudos[1] == conteudos[2] == conteudos[3]
    with zipfile.ZipFile(tmp_path / "ir.epub") as arquivo:
        assert "OEBPS/content.opf" in arquivo.namelist()
        assert "OEBPS/content.xhtml" not in arquivo.namelist()


def test_a_fachada_aplica_a_revisao_uma_vez_sobre_as_paginas_que_recebe(tmp_path):
    """A janela passa as páginas cruas do seu extrator e a fachada aplica a
    revisão: o bloco rejeitado sai, o texto revisado entra, e as páginas
    passadas ficam como estavam."""
    from core.editorial_review import ReviewSession

    paginas = [_pagina()]
    documento = paginas_extraidas_para_documento(paginas, document_id="livro", title="Livro")
    sessao = ReviewSession(documento)
    sessao.edit("block-livro-p0001-b0001", "Black to play and draw.")
    sessao.reject("block-livro-p0001-b0003")
    relatorio = EditorialPipeline().export(
        sessao.document, tmp_path / "revisado.epub",
        OpcoesDaFachada(format="epub", paginas=paginas,
                        escritor={"titulo": "Livro", "autor": "Ninguém"}))
    assert relatorio.metadata["paginas"] == "leitor"
    xhtml = _xhtml(tmp_path / "revisado.epub")
    assert "Black to play and draw." in xhtml and "The king walks." not in xhtml
    assert [type(b).__name__ for b in paginas[0].blocos] == [
        "Paragrafo", "Paragrafo", "Figura", "Paragrafo"]
    assert paginas[0].blocos[1].texto == "White to play and win."
    with zipfile.ZipFile(tmp_path / "revisado.epub") as arquivo:
        opf = arquivo.read("OEBPS/content.opf").decode("utf-8")
    assert "<dc:creator>Ninguém</dc:creator>" in opf
