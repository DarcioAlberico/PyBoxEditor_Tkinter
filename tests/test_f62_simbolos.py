"""
Testes da F62 — o recorte da fonte de símbolos.

O `gerar_fonte_de_simbolos.py` roda à mão e o produto dele é versionado, o que
cria um jeito silencioso de o livro sair errado: **o alfabeto do modelo cresce e
o recorte não**. Aí o símbolo novo vira quadradinho no EPUB, e nada acusa.

É isso que os testes abaixo prendem, e é por isso que eles leem o
`model_meta.json` em vez de uma lista fixa: a fonte de verdade sobre o que o
livro pode conter é o modelo, não este arquivo.

Rodar sem pytest:      python tests/test_f62_simbolos.py
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from core import exportar
from gerar_fonte_de_simbolos import (ARCO, DESENHADOS, EMPRESTADOS, ORIGEM,
                                     cobertura, simbolos_do_modelo)


def test_o_recorte_cobre_o_que_a_fonte_inteira_cobre():
    """
    Se o modelo aprender um símbolo que só a fonte inteira desenha, este teste
    acusa e o remédio é uma linha: `python gerar_fonte_de_simbolos.py`.
    """
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    simbolos = simbolos_do_modelo()
    inteira = cobertura(ORIGEM, simbolos)
    recorte = cobertura(exportar.SUBSET_DOS_SIMBOLOS, simbolos)
    faltando = [c for c in inteira if c not in recorte]

    assert not faltando, (
        f"o recorte não desenha {' '.join(faltando)} — rode "
        f"`python gerar_fonte_de_simbolos.py`")


def test_o_recorte_e_muito_menor_que_a_fonte_inteira():
    """A razão de ele existir. 641 KB para desenhar quarenta e um glifos."""
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    assert (os.path.getsize(exportar.SUBSET_DOS_SIMBOLOS)
            < os.path.getsize(ORIGEM) * 0.05)


def test_as_figurinas_estao_no_recorte():
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")
    assert cobertura(exportar.SUBSET_DOS_SIMBOLOS, "♔♕♖♗♘♙") == "♔♕♖♗♘♙"


def test_os_emprestados_estao_no_recorte():
    """
    O outro jeito de o livro sair errado, e o primeiro teste não pega este.

    Ele compara o recorte com a fonte inteira, e os emprestados **não estão na
    fonte inteira** — é por isso que foram emprestados. Some um par do
    `EMPRESTADOS` e a comparação com a Noto continua verde, com o `→` virando
    quadradinho no EPUB. Quem cobra estes é esta lista.
    """
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    pedidos = "".join(EMPRESTADOS)
    tem = cobertura(exportar.SUBSET_DOS_SIMBOLOS, pedidos)
    faltando = [c for c in pedidos if c not in tem]

    assert not faltando, (
        f"o recorte não desenha {' '.join(faltando)} — rode "
        f"`python gerar_fonte_de_simbolos.py`")


def test_os_desenhados_estao_no_recorte():
    """
    A terceira maneira de o livro sair errado, e as outras duas não pegam esta.

    O `⌓` não está na Noto nem vem de fonte nenhuma: o contorno dele é escrito
    pelo próprio script. Tirar a chamada de `desenhar` do `main` deixaria os
    dois testes acima verdes e o `$142` quadradinho na página.
    """
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    pedidos = "".join(DESENHADOS)
    tem = cobertura(exportar.SUBSET_DOS_SIMBOLOS, pedidos)
    faltando = [c for c in pedidos if c not in tem]

    assert not faltando, (
        f"o recorte não desenha {' '.join(faltando)} — rode "
        f"`python gerar_fonte_de_simbolos.py`")


def test_o_arco_do_recorte_tem_a_proporcao_do_livro():
    """
    O motivo de o `⌓` ser desenhado em vez de emprestado, preso num número.

    Ter glifo não basta: o `b` da SkakNew **tinha** glifo, e era uma meia-elipse
    de 0,53 onde as 64 amostras de `training_data/sym_8979` medem 0,714. Um
    empréstimo que voltasse a entrar aqui passaria nos testes de cobertura e
    devolveria o arco errado à página — é este número que o impede.
    """
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    from fontTools.pens.boundsPen import BoundsPen
    from fontTools.ttLib import TTFont

    fonte = TTFont(exportar.SUBSET_DOS_SIMBOLOS)
    glifos = fonte.getGlyphSet()
    nome = fonte.getBestCmap()[ord("⌓")]
    caneta = BoundsPen(glifos)
    glifos[nome].draw(caneta)
    x0, y0, x1, y1 = caneta.bounds

    proporcao = (y1 - y0) / (x1 - x0)
    alvo = ARCO["altura"] / ARCO["largura"]
    assert abs(proporcao - alvo) < 0.01, (
        f"o arco está em {proporcao:.3f} de altura sobre largura, e o livro "
        f"imprime {alvo:.3f}")
    assert y0 == 0, "o arco assenta na linha de base"


def test_o_cruz_do_xeque_nao_e_pedido_a_fonte():
    """
    O alfabeto do modelo tem o `✝` e **nenhum arquivo pode contê-lo**.

    São 2.901 amostras — a cruz cheia que estes livros imprimem no xeque —, e
    mesmo assim a F68 pôs o `notacao.SINONIMOS_DE_SAIDA` nos três lugares em que
    caractere de modelo vira texto de arquivo: ele sai `+` no livro, no PGN e na
    camada do PDF. Pedir glifo para ele seria pedir desenho para o que ninguém
    alcança, e o `ainda sem glifo` cobraria uma dívida paga para sempre.

    Se um dia o sinônimo cair, este teste cai junto — e aí o glifo passa a ser
    preciso de verdade.
    """
    from core import notacao

    assert "✝" in notacao.SINONIMOS_DE_SAIDA
    assert "✝" not in simbolos_do_modelo()


def test_o_recorte_e_o_escolhido_quando_basta():
    """Entre duas fontes que cobrem o mesmo, ganha a que pesa 1,2% da outra."""
    escolha = exportar.fonte_dos_simbolos("1.e4 ♖xf3 ♕d5")
    assert escolha
    if os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        assert escolha[1] == exportar.SUBSET_DOS_SIMBOLOS


def test_sem_recorte_a_fonte_inteira_assume():
    """
    O recorte é produto de script; ele pode não ter sido gerado num clone novo.
    Nesse caso o livro sai maior, e não sai errado.
    """
    original = exportar.SUBSET_DOS_SIMBOLOS
    exportar.SUBSET_DOS_SIMBOLOS = os.path.join(os.path.dirname(original),
                                                "nao-existe.ttf")
    try:
        escolha = exportar.fonte_dos_simbolos("1.e4 ♖xf3")
        assert escolha and escolha[1] == ORIGEM
    finally:
        exportar.SUBSET_DOS_SIMBOLOS = original


def test_a_familia_do_recorte_nao_se_passa_pela_original():
    """
    Subset é modificação, e a OFL pede que a modificada não use o nome da
    original. Há um motivo prático junto: uma família com o nome da Noto
    instalada na máquina de quem abre o arquivo brigaria com esta, que tem
    quarenta e um glifos.
    """
    if not os.path.exists(exportar.SUBSET_DOS_SIMBOLOS):
        pytest.skip("recorte não gerado (rode gerar_fonte_de_simbolos.py)")

    familia = exportar._familia(exportar.SUBSET_DOS_SIMBOLOS)
    assert "Noto" not in familia
    assert familia == "Simbolos de Xadrez"


# ----------------------------------------------------------------------
# O que estava errado no arquivo e ninguém via
# ----------------------------------------------------------------------

def test_o_idioma_do_livro_nao_e_o_do_programa():
    """
    Era `pt` fixo, e estes livros são em inglês — a mesma distinção que a §5.8
    da SPEC faz para o léxico. Livro declarado no idioma errado é hifenização
    errada e leitor de tela lendo notação inglesa com fonemas portugueses.
    """
    import tempfile
    import zipfile

    from core import livro

    paginas = [livro.PaginaExtraida(numero=0,
                                    blocos=[livro.Paragrafo("The rook is bad.")])]
    with tempfile.TemporaryDirectory() as tmp:
        caminho = exportar.para_epub(paginas, os.path.join(tmp, "a.epub"))
        with zipfile.ZipFile(caminho) as z:
            opf = z.read("OEBPS/content.opf").decode("utf-8")
            pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")

        assert "<dc:language>en</dc:language>" in opf
        assert 'xml:lang="en"' in pagina

        outro = exportar.para_epub(paginas, os.path.join(tmp, "b.epub"),
                                   idioma="pt")
        with zipfile.ZipFile(outro) as z:
            assert "<dc:language>pt</dc:language>" in z.read(
                "OEBPS/content.opf").decode("utf-8")


def test_o_titulo_do_paragrafo_deixa_de_ser_letra_morta():
    """
    O campo existe na `Paragrafo` desde a F2.6 e os dois exportadores o
    ignoravam. Nada o marca ainda — isso é trabalho de quem detectar título na
    página —, mas quem marcar agora encontra os dois formatos prontos.
    """
    import tempfile
    import zipfile

    from docx import Document

    from core import livro

    paginas = [livro.PaginaExtraida(
        numero=0, blocos=[livro.Paragrafo("Chapter 1", titulo=True),
                          livro.Paragrafo("Prosa comum.")])]
    with tempfile.TemporaryDirectory() as tmp:
        with zipfile.ZipFile(exportar.para_epub(
                paginas, os.path.join(tmp, "a.epub"))) as z:
            pagina = z.read("OEBPS/pagina-0001.xhtml").decode("utf-8")
        assert '<h2 id="t1-1">Chapter 1</h2>' in pagina
        assert "<p>Prosa comum.</p>" in pagina or (
            '<p class="primeira">Prosa comum.</p>' in pagina)

        doc = Document(exportar.para_docx(paginas, os.path.join(tmp, "a.docx")))
        assert doc.paragraphs[0].style.name == "Heading 2"
        assert doc.paragraphs[1].style.name != "Heading 2"


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
